import frappe

@frappe.whitelist()
def check_return_accounts(return_invoice="RBJM26H040001"):
    ret = frappe.get_doc("POS Invoice", return_invoice)
    orig = frappe.get_doc("POS Invoice", ret.return_against)

    print("=== RETURN ===")
    print(ret.name)
    print(ret.debit_to)

    print()
    print("=== ORIGINAL ===")
    print(orig.name)
    print(orig.debit_to)

    return {
        "return_debit_to": ret.debit_to,
        "original_debit_to": orig.debit_to
    }


@frappe.whitelist()
def test_pos_return_real_submit(pos_invoice="RBJM26H040001"):
    original = frappe.get_doc("POS Invoice", pos_invoice)

    print("=== ORIGINAL ===")
    print("Name:", original.name)
    print("Company:", original.company)
    print("Is Return:", original.is_return)
    print("Return Against:", original.return_against)
    print()

    try:
        # copy dokumen supaya tidak mengubah dokumen asli
        doc = frappe.copy_doc(original)
        doc.name = None
        doc.docstatus = 0

        print("=== INSERT TEST DOC ===")
        doc.insert(ignore_permissions=True)
        print("Inserted:", doc.name)

        print()
        print("=== SUBMIT TEST DOC ===")

        doc.submit()

        print("SUBMIT SUCCESS")

        # rollback supaya tidak tersimpan
        frappe.db.rollback()

        return "SUBMIT SUCCESS"

    except Exception:
        print()
        print("=== REAL ERROR ===")
        print(frappe.get_traceback())

        frappe.db.rollback()

        return "SUBMIT FAILED"

@frappe.whitelist()
def test_pos_return_coa(pos_invoice="RBJM26H040001"):
    doc = frappe.get_doc("POS Invoice", pos_invoice)

    print("=== POS INVOICE DEBUG ===")
    print("Name       :", doc.name)
    print("Company    :", doc.company)
    print("Is Return  :", doc.is_return)
    print("Return Agst:", doc.return_against)
    print()

    print("=== ITEMS ===")
    for row in doc.items:
        acc_company = None

        if row.income_account:
            acc_company = frappe.db.get_value(
                "Account",
                row.income_account,
                "company"
            )

        status = "OK" if acc_company == doc.company else "MISMATCH"

        print(
            f"ITEM={row.item_code} | ACCOUNT={row.income_account} | "
            f"ACC_COMPANY={acc_company} | STATUS={status}"
        )

    print()
    print("=== TAXES ===")
    for row in doc.taxes:
        acc_company = None

        if row.account_head:
            acc_company = frappe.db.get_value(
                "Account",
                row.account_head,
                "company"
            )

        status = "OK" if acc_company == doc.company else "MISMATCH"

        print(
            f"TAX={row.description} | ACCOUNT={row.account_head} | "
            f"ACC_COMPANY={acc_company} | STATUS={status}"
        )

    print()
    print("=== POS PROFILE ===")
    if doc.pos_profile:
        profile_company = frappe.db.get_value(
            "POS Profile",
            doc.pos_profile,
            "company"
        )

        status = "OK" if profile_company == doc.company else "MISMATCH"

        print(
            f"POS_PROFILE={doc.pos_profile} | PROFILE_COMPANY={profile_company} | STATUS={status}"
        )

    print()
    print("=== COST CENTER ===")
    if doc.cost_center:
        cc_company = frappe.db.get_value(
            "Cost Center",
            doc.cost_center,
            "company"
        )

        status = "OK" if cc_company == doc.company else "MISMATCH"

        print(
            f"COST_CENTER={doc.cost_center} | CC_COMPANY={cc_company} | STATUS={status}"
        )

    print()
    print("=== VALIDATE GL ===")

    try:
        gl_entries = doc.get_gl_entries()

        for gle in gl_entries:
            acc_company = frappe.db.get_value(
                "Account",
                gle.account,
                "company"
            )

            status = "OK" if acc_company == doc.company else "MISMATCH"

            print(
                f"GL ACCOUNT={gle.account} | DR={gle.debit} | CR={gle.credit} | "
                f"ACC_COMPANY={acc_company} | STATUS={status}"
            )

    except Exception as e:
        print("GL ERROR:", frappe.get_traceback())

    return "Done"

def test_pos_submit():
    pos = frappe.get_doc("POS Invoice", "BJB26H040004")

    print("docstatus:", pos.docstatus)

    try:
        pos.submit()
        print("SUBMIT OK")
    except Exception:
        print(frappe.get_traceback())

@frappe.whitelist()
def repair_gl_entry():	
	doctype = "Sales Invoice"
	docname = "ACC-SINV-2025-00063-1"

	docu = frappe.get_doc(doctype, docname)	
	delete_gl = frappe.db.sql(""" DELETE FROM `tabGL Entry` WHERE voucher_no = "{}" """.format(docname))
	docu.make_gl_entries()


@frappe.whitelist()
def create_tutup_kasir():
	from erpnext.accounts.doctype.pos_closing_entry.pos_closing_entry import make_closing_entry_from_opening
	closing_entry = make_closing_entry_from_opening(frappe.get_doc("POS Opening Entry","POS-OPE-2025-00009"))
	closing_entry.save()
	
def start_import():
	doc = frappe.get_doc("Data Import","Customer Import on 2026-09-15 10:22:43.693006")
	doc.start_import()

def start_import2():
	doc = frappe.get_doc("Data Import","Item Price Import on 2026-01-01 05:23:08.379058")
	doc.start_import()
def rename_customer():
	data = frappe.db.sql("select name , custom_kode from `tabCustomer` where name != custom_kode", as_list=1)
	count=0
	for row in data:
		frappe.rename_doc("Customer",row[0],row[1])
		frappe.db.commit()
		count=count+1
		print(count)


def debug():

    print("\n========== PI ==========")

    pi = frappe.get_doc("Purchase Invoice", "PI-BJM4260900005")

    print("name               :", pi.name)
    print("status             :", pi.status)
    print("docstatus          :", pi.docstatus)
    print("company            :", pi.company)
    print("supplier           :", pi.supplier)
    print("credit_to          :", pi.credit_to)
    print("grand_total        :", pi.grand_total)
    print("outstanding_amount :", pi.outstanding_amount)
    print("is_paid            :", pi.is_paid)


    print("\n========== PI GL ==========")

    pi_gl = frappe.db.sql("""
        SELECT
            account,
            party_type,
            party,
            debit,
            credit,
            debit_in_account_currency,
            credit_in_account_currency,
            against,
            voucher_type,
            voucher_no,
            is_cancelled
        FROM `tabGL Entry`
        WHERE voucher_type = 'Purchase Invoice'
        AND voucher_no = %s
        AND is_cancelled = 0
        ORDER BY creation
    """, pi.name, as_dict=True)

    for row in pi_gl:
        print(row)


    print("\n========== PAYMENT ENTRY ==========")

    pe = frappe.get_doc("Payment Entry", "BJMPE26-090002-2")

    print("name             :", pe.name)
    print("docstatus        :", pe.docstatus)
    print("payment_type     :", pe.payment_type)
    print("party_type       :", pe.party_type)
    print("party            :", pe.party)
    print("company          :", pe.company)
    print("paid_from        :", pe.paid_from)
    print("paid_to          :", pe.paid_to)
    print("paid_amount      :", pe.paid_amount)
    print("received_amount  :", pe.received_amount)


    print("\n========== PE REFERENCES ==========")

    for r in pe.references:
        print({
            "doctype": r.reference_doctype,
            "name": r.reference_name,
            "allocated": r.allocated_amount,
            "total": r.total_amount,
            "outstanding": r.outstanding_amount
        })


    print("\n========== PAYMENT GL ==========")

    pe_gl = frappe.db.sql("""
        SELECT
            account,
            party_type,
            party,
            debit,
            credit,
            debit_in_account_currency,
            credit_in_account_currency,
            against,
            voucher_type,
            voucher_no,
            is_cancelled
        FROM `tabGL Entry`
        WHERE voucher_type = 'Payment Entry'
        AND voucher_no = %s
        AND is_cancelled = 0
        ORDER BY creation
    """, pe.name, as_dict=True)

    for row in pe_gl:
        print(row)


    print("\n========== SUPPLIER LEDGER ==========")

    ledger = frappe.db.sql("""
        SELECT
            posting_date,
            account,
            debit,
            credit,
            voucher_type,
            voucher_no,
            against_voucher_type,
            against_voucher
        FROM `tabGL Entry`
        WHERE party_type = 'Supplier'
        AND party = %s
        AND is_cancelled = 0
        AND company = %s
        ORDER BY posting_date, creation
    """, pi.supplier, pi.company, as_dict=True)

    for row in ledger:
        print(row)


    print("\n========== PI PAYMENT STATUS CALCULATION ==========")

    print("PI outstanding_amount :", frappe.db.get_value(
        "Purchase Invoice",
        pi.name,
        "outstanding_amount"
    ))

    print("PI status             :", frappe.db.get_value(
        "Purchase Invoice",
        pi.name,
        "status"
    ))

    print("PE allocated_amount   :", frappe.db.get_value(
        "Payment Entry Reference",
        {
            "parent": pe.name,
            "reference_doctype": "Purchase Invoice",
            "reference_name": pi.name
        },
        "allocated_amount"
    ))


def upload_ste(docname="Stock Entry Import on 2026-09-02 16:22:56.860458"):
    #  Stock Entry Import on 2026-09-02 12:18:33.009419
    #  Stock Entry Import on 2026-09-02 13:08:27.484219
    #  Stock Entry Import on 2026-09-02 13:30:58.977945
    #  Stock Entry Import on 2026-09-02 13:57:18.287758
    #  Stock Entry Import on 2026-09-02 14:12:51.310292

    # Stock Entry Import on 2026-09-02 14:40:25.224303
    # Stock Entry Import on 2026-09-02 15:05:57.972782
    # Stock Entry Import on 2026-09-02 15:29:12.925658
    # Stock Entry Import on 2026-09-02 15:57:07.834753
    # Stock Entry Import on 2026-09-02 16:22:56.860458
    doc = frappe.get_doc("Data Import", docname)
    doc.start_import()


import os
import openpyxl


def split_import_backend_alan(chunk_size=50000):
    current_dir = os.path.dirname(os.path.abspath(__file__))

    input_file = os.path.join(
        current_dir,
        "Item Price Import per 31 ags 2026.xlsx"
    )

    output_dir = current_dir

    wb = openpyxl.load_workbook(input_file, read_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header = rows[0]
    data = rows[1:]

    total = len(data)
    total_files = (total + chunk_size - 1) // chunk_size

    print(f"Total data : {total}")
    print(f"Ukuran/file: {chunk_size}")
    print(f"Total file : {total_files}")

    for i in range(total_files):
        start = i * chunk_size
        end = min(start + chunk_size, total)

        output_file = os.path.join(
            output_dir,
            f"Item Price Import per 31 ags 2026 {i + 1}.xlsx"
        )

        new_wb = openpyxl.Workbook()
        new_ws = new_wb.active

        new_ws.append(header)

        for row in data[start:end]:
            new_ws.append(row)

        new_wb.save(output_file)
        new_wb.close()

        print(
            f"File {i + 1}: "
            f"{end - start} data "
            f"({start + 1}-{end}) -> {output_file}"
        )

    print("Selesai.")


def hapus_supplier_dari_excel():
    current_dir = os.path.dirname(os.path.abspath(__file__))
    file_path = os.path.join(current_dir, "Hapus Supplier.xlsx")

    wb = openpyxl.load_workbook(
        file_path,
        read_only=True,
        data_only=True
    )

    ws = wb.active

    # Cari kolom ID
    id_col = None

    for col in range(1, ws.max_column + 1):
        header = ws.cell(1, col).value

        if header and str(header).strip() == "ID":
            id_col = col
            break

    if not id_col:
        frappe.throw("Kolom ID tidak ditemukan")

    deleted = 0
    not_found = 0
    failed = 0

    for row in range(2, ws.max_row + 1):
        supplier_id = ws.cell(row, id_col).value

        if not supplier_id:
            continue

        supplier_id = str(supplier_id).strip()

        if not frappe.db.exists("Supplier", supplier_id):
            print(f"[NOT FOUND] {supplier_id}")
            not_found += 1
            continue

        supplier_name = frappe.db.get_value(
            "Supplier",
            supplier_id,
            "supplier_name"
        )

        try:
            frappe.delete_doc(
                "Supplier",
                supplier_id,
                force=True
            )

            print(f"[DELETED] {supplier_id} - {supplier_name}")
            deleted += 1

        except Exception as e:
            print(f"[FAILED] {supplier_id} - {supplier_name} - {e}")
            failed += 1

    wb.close()

    frappe.db.commit()

    print("")
    print("========== HASIL ==========")
    print(f"Deleted   : {deleted}")
    print(f"Not Found : {not_found}")
    print(f"Failed    : {failed}")


from difflib import get_close_matches

def check_and_fix_uom_excel(
    input_file="Item Price Import per 31 ags 2026.xlsx",
    uom_column="UOM"
):
    """
    Cek UOM di Excel terhadap UOM yang ada di ERPNext.

    - Beda huruf besar/kecil -> otomatis diperbaiki
    - Typo ringan -> otomatis diperbaiki
    - UOM yang tidak ditemukan -> tidak diubah dan ditampilkan di rekap
    - Hasil disimpan ke file _fixed.xlsx
    """

    # ============================================================
    # FILE EXCEL DI FOLDER YANG SAMA DENGAN FILE PY
    # ============================================================

    current_dir = os.path.dirname(
        os.path.abspath(__file__)
    )

    input_file = os.path.join(
        current_dir,
        input_file
    )

    if not os.path.exists(input_file):
        frappe.throw(
            f"File Excel tidak ditemukan:\n{input_file}"
        )

    # ============================================================
    # AMBIL UOM DARI ERPNext
    # ============================================================

    system_uoms = frappe.get_all(
        "UOM",
        pluck="name"
    )

    uom_map = {
        uom.strip().lower(): uom
        for uom in system_uoms
        if uom
    }

    # ============================================================
    # BUKA EXCEL
    # ============================================================

    wb = openpyxl.load_workbook(input_file)
    ws = wb.active

    # ============================================================
    # CARI KOLOM UOM
    # ============================================================

    uom_col = None

    for cell in ws[1]:

        if (
            cell.value is not None
            and str(cell.value).strip().lower() == uom_column.lower()
        ):
            uom_col = cell.column
            break

    if uom_col is None:
        frappe.throw(
            f"Kolom '{uom_column}' tidak ditemukan di Excel."
        )

    # ============================================================
    # REKAP UOM YANG TIDAK DITEMUKAN
    # ============================================================

    not_found = set()

    # ============================================================
    # CEK SETIAP ROW
    # ============================================================

    for row_number in range(2, ws.max_row + 1):

        cell = ws.cell(
            row=row_number,
            column=uom_col
        )

        if cell.value is None:
            continue

        original = str(cell.value).strip()

        if not original:
            continue

        normalized = original.lower()

        # ========================================================
        # 1. UOM ADA, TERMASUK BEDA BESAR/KECIL
        # ========================================================

        if normalized in uom_map:

            correct_uom = uom_map[normalized]

            if original != correct_uom:
                cell.value = correct_uom

            continue

        # ========================================================
        # 2. COBA TYPO RINGAN
        # ========================================================

        matches = get_close_matches(
            normalized,
            list(uom_map.keys()),
            n=1,
            cutoff=0.85
        )

        if matches:

            matched = matches[0]
            correct_uom = uom_map[matched]

            cell.value = correct_uom

        else:

            # ====================================================
            # BENAR-BENAR TIDAK ADA
            # ====================================================

            not_found.add(original)

    # ============================================================
    # SIMPAN KE FILE _FIXED
    # ============================================================

    output_file = os.path.join(
        current_dir,
        "Item Price Import per 31 ags 2026_fixed.xlsx"
    )

    wb.save(output_file)

    # ============================================================
    # REKAP SINGKAT
    # ============================================================

    print()
    print("=" * 50)
    print("UOM CHECK SELESAI")
    print("=" * 50)

    if not_found:

        print("UOM YANG TIDAK ADA DI SYSTEM:")

        for uom in sorted(not_found):
            print(f"- {uom}")

    else:

        print("Semua UOM ditemukan di system.")

    print()
    print(f"Total UOM tidak ditemukan: {len(not_found)}")
    print(f"File hasil: {output_file}")

    return {
        "not_found": sorted(not_found),
        "output_file": output_file,
    }


def patch_item_price_item_name():
    print("=" * 70)
    print("PATCH ITEM PRICE - ITEM NAME")
    print("=" * 70)

    # Cek jumlah yang perlu diperbaiki
    before = frappe.db.sql("""
        SELECT COUNT(*)
        FROM `tabItem Price` ip
        INNER JOIN `tabItem` i
            ON i.name = ip.item_code
        WHERE COALESCE(ip.item_name, '') = ''
    """)[0][0]

    print(f"[CHECK] Item Price tanpa item_name : {before}")

    if not before:
        print("[DONE] Tidak ada data yang perlu dipatch.")
        return

    # Bulk UPDATE
    frappe.db.sql("""
        UPDATE `tabItem Price` ip
        INNER JOIN `tabItem` i
            ON i.name = ip.item_code
        SET
            ip.item_name = i.item_name,
            ip.modified = NOW(),
            ip.modified_by = %(modified_by)s
        WHERE COALESCE(ip.item_name, '') = ''
    """, {
        "modified_by": frappe.session.user
    })

    frappe.db.commit()

    # Verifikasi
    after = frappe.db.sql("""
        SELECT COUNT(*)
        FROM `tabItem Price` ip
        INNER JOIN `tabItem` i
            ON i.name = ip.item_code
        WHERE COALESCE(ip.item_name, '') = ''
    """)[0][0]

    patched = before - after

    print("-" * 70)
    print(f"[PATCHED] {patched}")
    print(f"[REMAIN]  {after}")
    print("[COMMIT] Berhasil")
    print("=" * 70)



import os
import openpyxl
import frappe


def merge_item_price_excel(
    input_file="Item Price (Import per 31 ags 2026) (2).xlsx",
    output_file=None,
):
    """
    Merge seluruh sheet Excel menjadi satu sheet.

    Ketentuan:
    - Semua sheet dianggap memiliki struktur kolom yang sama.
    - Baris pertama setiap sheet dianggap sebagai header.
    - Header hanya ditulis satu kali.
    - Data dari semua sheet digabungkan.
    - Ditambahkan kolom `source_sheet` untuk mengetahui asal sheet.
    """

    # ==========================================================
    # 1. FILE PATH
    # ==========================================================

    current_dir = os.path.dirname(
        os.path.abspath(__file__)
    )

    if not os.path.isabs(input_file):
        input_file = os.path.join(
            current_dir,
            input_file
        )

    if not os.path.exists(input_file):
        frappe.throw(
            f"File Excel tidak ditemukan:\n{input_file}"
        )

    # ==========================================================
    # 2. OUTPUT FILE
    # ==========================================================

    if output_file is None:

        base_name = os.path.splitext(
            os.path.basename(input_file)
        )[0]

        output_file = os.path.join(
            current_dir,
            f"{base_name}_merged.xlsx"
        )

    elif not os.path.isabs(output_file):

        output_file = os.path.join(
            current_dir,
            output_file
        )

    # ==========================================================
    # 3. LOAD EXCEL
    # ==========================================================

    wb = openpyxl.load_workbook(
        input_file,
        read_only=True,
        data_only=True,
    )

    sheet_names = wb.sheetnames

    if not sheet_names:

        wb.close()

        frappe.throw(
            "Excel tidak memiliki sheet."
        )

    print("=" * 70)
    print("MERGE ITEM PRICE EXCEL")
    print("=" * 70)
    print(f"File       : {input_file}")
    print(f"Jumlah tab : {len(sheet_names)}")
    print()

    # ==========================================================
    # 4. BUAT WORKBOOK BARU
    # ==========================================================

    output_wb = openpyxl.Workbook()

    output_ws = output_wb.active

    output_ws.title = "Item Price"

    # ==========================================================
    # 5. HEADER
    # ==========================================================

    first_sheet = wb[sheet_names[0]]

    headers = [
        cell.value
        for cell in first_sheet[1]
    ]

    # Tambahkan source sheet
    output_ws.append([
        *headers,
        "source_sheet",
    ])

    # ==========================================================
    # 6. MERGE SEMUA SHEET
    # ==========================================================

    total_rows = 0

    for sheet_name in sheet_names:

        ws = wb[sheet_name]

        print(
            f"[PROCESS] Sheet: {sheet_name}"
        )

        sheet_rows = 0

        # ------------------------------------------------------
        # Mulai dari row 2 karena row 1 adalah header
        # ------------------------------------------------------

        for row in ws.iter_rows(
            min_row=2,
            values_only=True
        ):

            # --------------------------------------------------
            # Skip baris kosong
            # --------------------------------------------------

            if not any(
                value is not None
                for value in row
            ):
                continue

            output_ws.append([
                *row,
                sheet_name,
            ])

            sheet_rows += 1
            total_rows += 1

        print(
            f"          {sheet_rows} row"
        )

    # ==========================================================
    # 7. SAVE
    # ==========================================================

    output_wb.save(
        output_file
    )

    output_wb.close()
    wb.close()

    # ==========================================================
    # 8. SUMMARY
    # ==========================================================

    print()
    print("=" * 70)
    print("MERGE SELESAI")
    print("=" * 70)

    print(
        f"Sheet digabung : {len(sheet_names)}"
    )

    print(
        f"Total row      : {total_rows}"
    )

    print(
        f"Output         : {output_file}"
    )

    print("=" * 70)

    return {
        "input_file": input_file,
        "output_file": output_file,
        "sheet_names": sheet_names,
        "total_rows": total_rows,
    }


import os
import openpyxl
import frappe

from datetime import datetime, date
def import_item_price_from_excel(
    input_file="insert item price bjb.xlsx",
    commit=True,
    batch_size=1000
):
    import os
    from datetime import datetime, date

    import openpyxl
    import frappe

    # ============================================================
    # PATH
    # ============================================================
    file_path = os.path.join(
        os.path.dirname(__file__),
        input_file
    )

    if not os.path.exists(file_path):
        frappe.throw(f"File tidak ditemukan: {file_path}")

    base_name = os.path.splitext(input_file)[0]

    checkpoint_file = os.path.join(
        os.path.dirname(__file__),
        f"{base_name}_checkpoint.txt"
    )

    failed_file = os.path.join(
        os.path.dirname(__file__),
        f"{base_name}_not_imported.xlsx"
    )

    # ============================================================
    # HELPER
    # ============================================================
    def normalize_item_code(value):
        if value is None:
            return ""

        value = str(value).strip()

        # Excel bisa menyimpan apostrophe sebagai karakter awal
        if value.startswith("'"):
            value = value[1:].strip()

        return value

    def normalize_price(value):
        if value is None:
            return None

        # Jika Excel sudah membaca sebagai angka
        if isinstance(value, (int, float)):
            return frappe.utils.flt(value)

        value = str(value).strip()

        if not value:
            return None

        # Hilangkan format currency
        value = value.replace("Rp", "")
        value = value.replace("rp", "")
        value = value.replace(" ", "")

        # Contoh:
        # 94,909.00 -> 94909.00
        # 94909,50  -> 94909.50
        if "," in value and "." in value:
            value = value.replace(",", "")

        elif "," in value:
            value = value.replace(",", ".")

        try:
            return frappe.utils.flt(value)
        except Exception:
            return None

    def normalize_date(value):
        if value is None:
            return None

        if isinstance(value, datetime):
            return value.date()

        if isinstance(value, date):
            return value

        value = str(value).strip()

        if not value:
            return None

        formats = [
            "%d-%m-%Y",
            "%Y-%m-%d",
            "%d/%m/%Y",
            "%Y/%m/%d",
        ]

        for fmt in formats:
            try:
                return datetime.strptime(value, fmt).date()
            except ValueError:
                continue

        return None

    def normalize_bool(value):
        if value is None:
            return 0

        if isinstance(value, bool):
            return int(value)

        value = str(value).strip().lower()

        if value in ("1", "true", "yes", "y"):
            return 1

        return 0

    # ============================================================
    # LOAD EXCEL
    # ============================================================
    print("=" * 80)
    print("ITEM PRICE IMPORT START")
    print("=" * 80)
    print(f"File       : {input_file}")
    print(f"Batch size : {batch_size}")
    print(f"Commit     : {commit}")

    wb = openpyxl.load_workbook(
        file_path,
        read_only=True,
        data_only=True
    )

    ws = wb.active

    # ============================================================
    # HEADER
    # ============================================================
    headers = next(ws.iter_rows(values_only=True))

    header_map = {}

    for idx, header in enumerate(headers, start=1):
        if header is None:
            continue

        header_name = str(header).strip().lower()

        header_map[header_name] = idx

    def find_header(*names):
        for name in names:
            if name.lower() in header_map:
                return header_map[name.lower()]

        return None

    item_code_col = find_header(
        "item_code",
        "item code",
        "item"
    )

    uom_col = find_header(
        "uom",
        "stock_uom",
        "stock uom"
    )

    price_list_col = find_header(
        "price_list",
        "price list"
    )

    rate_col = find_header(
        "price_list_rate",
        "price list rate",
        "rate",
        "price"
    )

    currency_col = find_header(
        "currency"
    )

    buying_col = find_header(
        "buying"
    )

    selling_col = find_header(
        "selling"
    )

    valid_from_col = find_header(
        "valid_from",
        "valid from"
    )

    valid_upto_col = find_header(
        "valid_upto",
        "valid upto",
        "valid up to"
    )

    # ============================================================
    # REQUIRED HEADER VALIDATION
    # ============================================================
    missing_headers = []

    if not item_code_col:
        missing_headers.append("Item Code")

    if not uom_col:
        missing_headers.append("UOM")

    if not price_list_col:
        missing_headers.append("Price List")

    if not rate_col:
        missing_headers.append("Rate")

    if missing_headers:
        frappe.throw(
            "Header Excel tidak ditemukan: "
            + ", ".join(missing_headers)
        )

    print(f"Item Code  : kolom {item_code_col}")
    print(f"UOM        : kolom {uom_col}")
    print(f"Price List : kolom {price_list_col}")
    print(f"Rate       : kolom {rate_col}")

    # ============================================================
    # LOAD ALL UOM
    # ============================================================
    uom_rows = frappe.db.sql(
        """
        SELECT name
        FROM `tabUOM`
        """,
        as_dict=True
    )

    uom_map = {
        row.name.strip().lower(): row.name
        for row in uom_rows
    }

    print(f"UOM loaded : {len(uom_map)}")

    # ============================================================
    # CHECKPOINT
    # ============================================================
    start_row = 2

    if os.path.exists(checkpoint_file):
        try:
            with open(checkpoint_file, "r") as f:
                checkpoint = int(f.read().strip())

            start_row = checkpoint + 1

        except Exception:
            start_row = 2

    print(f"Start row  : {start_row}")

    if start_row > 2:
        print(
            f"Melanjutkan dari checkpoint row {start_row - 1}"
        )

    # ============================================================
    # FAILED ROWS
    # ============================================================
    failed_rows = []

    total_processed = 0
    total_imported = 0
    total_updated = 0
    total_failed = 0

    # ============================================================
    # ITERATE EXCEL
    # ============================================================
    batch = []
    excel_row_number = 1

    for row in ws.iter_rows(values_only=True):

        excel_row_number += 1

        if excel_row_number < start_row:
            continue

        batch.append(
            (
                excel_row_number,
                row
            )
        )

        if len(batch) >= batch_size:

            # ====================================================
            # PROCESS BATCH
            # ====================================================
            result = process_item_price_batch(
                batch=batch,
                item_code_col=item_code_col,
                uom_col=uom_col,
                price_list_col=price_list_col,
                rate_col=rate_col,
                currency_col=currency_col,
                buying_col=buying_col,
                selling_col=selling_col,
                valid_from_col=valid_from_col,
                valid_upto_col=valid_upto_col,
                uom_map=uom_map,
                normalize_item_code=normalize_item_code,
                normalize_price=normalize_price,
                normalize_date=normalize_date,
                normalize_bool=normalize_bool,
                commit=commit,
            )

            total_processed += result["processed"]
            total_imported += result["imported"]
            total_updated += result["updated"]
            total_failed += result["failed"]

            failed_rows.extend(
                result["failed_rows"]
            )

            # ====================================================
            # CHECKPOINT
            # ====================================================
            with open(checkpoint_file, "w") as f:
                f.write(str(excel_row_number))

            print(
                f"[BATCH] Row sampai {excel_row_number} | "
                f"Processed: {total_processed} | "
                f"Imported: {total_imported} | "
                f"Updated: {total_updated} | "
                f"Failed: {total_failed}"
            )

            batch = []

    # ============================================================
    # PROCESS LAST BATCH
    # ============================================================
    if batch:

        result = process_item_price_batch(
            batch=batch,
            item_code_col=item_code_col,
            uom_col=uom_col,
            price_list_col=price_list_col,
            rate_col=rate_col,
            currency_col=currency_col,
            buying_col=buying_col,
            selling_col=selling_col,
            valid_from_col=valid_from_col,
            valid_upto_col=valid_upto_col,
            uom_map=uom_map,
            normalize_item_code=normalize_item_code,
            normalize_price=normalize_price,
            normalize_date=normalize_date,
            normalize_bool=normalize_bool,
            commit=commit,
        )

        total_processed += result["processed"]
        total_imported += result["imported"]
        total_updated += result["updated"]
        total_failed += result["failed"]

        failed_rows.extend(
            result["failed_rows"]
        )

        # ========================================================
        # CHECKPOINT
        # ========================================================
        with open(checkpoint_file, "w") as f:
            f.write(str(batch[-1][0]))

        print(
            f"[LAST BATCH] Row sampai {batch[-1][0]} | "
            f"Processed: {total_processed} | "
            f"Imported: {total_imported} | "
            f"Updated: {total_updated} | "
            f"Failed: {total_failed}"
        )

    # ============================================================
    # SAVE FAILED EXCEL
    # ============================================================
    if failed_rows:

        failed_wb = openpyxl.Workbook()
        failed_ws = failed_wb.active

        # Header asli
        failed_ws.append(
            list(headers) + ["Error"]
        )

        for failed_row in failed_rows:

            failed_ws.append(
                list(failed_row["row"])
                + [failed_row["error"]]
            )

        failed_wb.save(
            failed_file
        )

    # ============================================================
    # CLOSE
    # ============================================================
    wb.close()

    # ============================================================
    # SUMMARY
    # ============================================================
    summary = {
        "file": input_file,
        "processed": total_processed,
        "imported": total_imported,
        "updated": total_updated,
        "failed": total_failed,
        "failed_file": failed_file if failed_rows else None,
        "checkpoint_file": checkpoint_file,
    }

    print("=" * 80)
    print("ITEM PRICE IMPORT FINISHED")
    print("=" * 80)
    print(f"Processed : {total_processed}")
    print(f"Imported  : {total_imported}")
    print(f"Updated   : {total_updated}")
    print(f"Failed    : {total_failed}")

    if failed_rows:
        print(f"Failed file: {failed_file}")

    print(f"Checkpoint : {checkpoint_file}")
    print("=" * 80)

    return summary


def process_item_price_batch(
    batch,
    item_code_col,
    uom_col,
    price_list_col,
    rate_col,
    currency_col,
    buying_col,
    selling_col,
    valid_from_col,
    valid_upto_col,
    uom_map,
    normalize_item_code,
    normalize_price,
    normalize_date,
    normalize_bool,
    commit=True,
):
    import frappe

    processed = 0
    imported = 0
    updated = 0
    failed = 0

    failed_rows = []

    # ============================================================
    # COLLECT ITEM CODE
    # ============================================================
    item_codes = set()

    for excel_row_number, row in batch:

        value = row[item_code_col - 1]

        item_code = normalize_item_code(value)

        if item_code:
            item_codes.add(item_code)

    # ============================================================
    # CHECK ITEMS
    # ============================================================
    existing_items = {}

    if item_codes:

        placeholders = ", ".join(
            ["%s"] * len(item_codes)
        )

        existing = frappe.db.sql(
            f"""
            SELECT
                name,
                item_name,
                description
            FROM `tabItem`
            WHERE name IN ({placeholders})
            """,
            tuple(item_codes),
            as_dict=True,
        )

        existing_items = {
            row["name"]: {
                "item_name": row["item_name"],
                "item_description": row["description"],
            }
            for row in existing
        }

    # ============================================================
    # VALID ROWS
    # ============================================================
    valid_rows = []

    for excel_row_number, row in batch:

        processed += 1

        try:
            # ====================================================
            # ITEM CODE
            # ====================================================
            item_code = normalize_item_code(
                row[item_code_col - 1]
            )

            if not item_code:
                raise Exception(
                    "Item Code kosong"
                )

            if item_code not in existing_items:
                raise Exception(
                    f"Item tidak ditemukan: {item_code}"
                )

            # ====================================================
            # ITEM MASTER DATA
            # ====================================================
            item_name = existing_items[item_code]["item_name"]
            item_description = existing_items[item_code]["item_description"]

            # ====================================================
            # UOM
            # ====================================================
            uom_value = row[uom_col - 1]

            if uom_value is None:
                raise Exception(
                    "UOM kosong"
                )

            uom_raw = str(
                uom_value
            ).strip()

            if not uom_raw:
                raise Exception(
                    "UOM kosong"
                )

            uom = uom_map.get(
                uom_raw.lower()
            )

            if not uom:
                raise Exception(
                    f"UOM tidak ditemukan: {uom_raw}"
                )

            # ====================================================
            # PRICE LIST
            # ====================================================
            price_list = row[
                price_list_col - 1
            ]

            if price_list is None:
                raise Exception(
                    "Price List kosong"
                )

            price_list = str(
                price_list
            ).strip()

            if not price_list:
                raise Exception(
                    "Price List kosong"
                )

            # ====================================================
            # RATE
            # ====================================================
            rate = normalize_price(
                row[rate_col - 1]
            )

            if rate is None:
                raise Exception(
                    f"Rate tidak valid: "
                    f"{row[rate_col - 1]}"
                )

            # ====================================================
            # CURRENCY
            # ====================================================
            currency = None

            if currency_col:
                currency = row[
                    currency_col - 1
                ]

                if currency is not None:
                    currency = str(
                        currency
                    ).strip()

            # ====================================================
            # BUYING
            # ====================================================
            buying = 0

            if buying_col:
                buying = normalize_bool(
                    row[buying_col - 1]
                )

            # ====================================================
            # SELLING
            # ====================================================
            selling = 0

            if selling_col:
                selling = normalize_bool(
                    row[selling_col - 1]
                )

            # ====================================================
            # VALID FROM
            # ====================================================
            valid_from = None

            if valid_from_col:

                valid_from = normalize_date(
                    row[valid_from_col - 1]
                )

                if row[valid_from_col - 1] is not None:

                    if valid_from is None:
                        raise Exception(
                            f"Valid From tidak valid: "
                            f"{row[valid_from_col - 1]}"
                        )

            # ====================================================
            # VALID UPTO
            # ====================================================
            valid_upto = None

            if valid_upto_col:

                valid_upto = normalize_date(
                    row[valid_upto_col - 1]
                )

                if row[valid_upto_col - 1] is not None:

                    if valid_upto is None:
                        raise Exception(
                            f"Valid Upto tidak valid: "
                            f"{row[valid_upto_col - 1]}"
                        )

            # ====================================================
            # VALID ROW
            # ====================================================
            valid_rows.append(
                {
                    "excel_row": excel_row_number,
                    "row": row,
                    "item_code": item_code,
                    "item_name": item_name,
                    "item_description": item_description,
                    "uom": uom,
                    "price_list": price_list,
                    "price_list_rate": rate,
                    "currency": currency,
                    "buying": buying,
                    "selling": selling,
                    "valid_from": valid_from,
                    "valid_upto": valid_upto,
                }
            )

        except Exception as e:

            failed += 1

            failed_rows.append(
                {
                    "row": row,
                    "error": str(e),
                }
            )

    # ============================================================
    # DEDUPLICATE
    # Last row wins
    # ============================================================
    unique_rows = {}

    for data in valid_rows:

        key = (
            data["item_code"],
            data["price_list"],
            data["uom"],
        )

        unique_rows[key] = data

    valid_rows = list(
        unique_rows.values()
    )

    # ============================================================
    # EXISTING ITEM PRICE
    # ============================================================
    existing_prices = {}

    if valid_rows:

        conditions = []
        values = []

        for data in valid_rows:

            conditions.append(
                """
                (
                    item_code = %s
                    AND price_list = %s
                    AND uom = %s
                )
                """
            )

            values.extend(
                [
                    data["item_code"],
                    data["price_list"],
                    data["uom"],
                ]
            )

        where_clause = " OR ".join(
            conditions
        )

        existing = frappe.db.sql(
            f"""
            SELECT
                name,
                item_code,
                price_list,
                uom
            FROM `tabItem Price`
            WHERE {where_clause}
            """,
            tuple(values),
            as_dict=True,
        )

        for row in existing:

            key = (
                row["item_code"],
                row["price_list"],
                row["uom"],
            )

            existing_prices[key] = row["name"]

    # ============================================================
    # INSERT / UPDATE
    # ============================================================
    for data in valid_rows:

        key = (
            data["item_code"],
            data["price_list"],
            data["uom"],
        )

        try:

            if key in existing_prices:

                # =================================================
                # UPDATE
                # =================================================
                name = existing_prices[key]

                frappe.db.sql(
                    """
                    UPDATE `tabItem Price`
                    SET
                        item_name = %s,
                        item_description = %s,
                        price_list_rate = %s,
                        currency = %s,
                        buying = %s,
                        selling = %s,
                        valid_from = %s,
                        valid_upto = %s,
                        modified = NOW(),
                        modified_by = %s
                    WHERE name = %s
                    """,
                    (
                        data["item_name"],
                        data["item_description"],
                        data["price_list_rate"],
                        data["currency"],
                        data["buying"],
                        data["selling"],
                        data["valid_from"],
                        data["valid_upto"],
                        frappe.session.user,
                        name,
                    ),
                )

                updated += 1

            else:

                # =================================================
                # INSERT
                # =================================================
                name = frappe.generate_hash(
                    length=10
                )

                frappe.db.sql(
                    """
                    INSERT INTO `tabItem Price`
                    (
                        name,
                        creation,
                        modified,
                        modified_by,
                        owner,
                        docstatus,
                        idx,
                        item_code,
                        item_name,
                        item_description,
                        price_list,
                        uom,
                        price_list_rate,
                        currency,
                        buying,
                        selling,
                        valid_from,
                        valid_upto
                    )
                    VALUES
                    (
                        %s,
                        NOW(),
                        NOW(),
                        %s,
                        %s,
                        0,
                        0,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s,
                        %s
                    )
                    """,
                    (
                        name,
                        frappe.session.user,
                        frappe.session.user,
                        data["item_code"],
                        data["item_name"],
                        data["item_description"],
                        data["price_list"],
                        data["uom"],
                        data["price_list_rate"],
                        data["currency"],
                        data["buying"],
                        data["selling"],
                        data["valid_from"],
                        data["valid_upto"],
                    ),
                )

                imported += 1

        except Exception as e:

            failed += 1

            failed_rows.append(
                {
                    "row": data["row"],
                    "error": str(e),
                }
            )

    # ============================================================
    # COMMIT
    # ============================================================
    if commit:
        frappe.db.commit()

    return {
        "processed": processed,
        "imported": imported,
        "updated": updated,
        "failed": failed,
        "failed_rows": failed_rows,
    }

def import_item_price_from_excel_no_uom(
    input_file="Item Price untuk Stock 0.xlsx",
    commit=True,
    batch_size=1000,
):
    """
    Import Item Price dari Excel secara batch menggunakan direct SQL.

    Format Excel:
    - ID
    - Item Code
    - Item Name
    - Price List
    - Buying
    - Selling
    - Rate
    - Valid From
    - Valid Upto (opsional)
    - Currency (opsional)

    Aturan:
    - UOM TIDAK diambil dari Excel.
    - UOM otomatis diambil dari Item.stock_uom.
    - Item Name otomatis diambil dari Item.item_name.
    - Item divalidasi dengan 1 query per batch.
    - Item Price existing diambil dengan 1 query per batch.
    - Insert / update menggunakan direct SQL.
    - Tidak menggunakan frappe.get_doc().insert().
    - Error validasi per row dicatat dan proses tetap lanjut.
    - Setiap batch di-commit.
    - Checkpoint ditulis setelah commit berhasil.
    - Failed rows disimpan ke *_not_imported.xlsx.
    - Jika kombinasi yang sama muncul beberapa kali dalam batch,
      row terakhir akan dipakai.
    """

    # ==========================================================
    # 1. FILE PATH
    # ==========================================================

    current_dir = os.path.dirname(
        os.path.abspath(__file__)
    )

    if not os.path.isabs(input_file):
        input_file = os.path.join(
            current_dir,
            input_file
        )

    if not os.path.exists(input_file):
        frappe.throw(
            f"File Excel tidak ditemukan:\n{input_file}"
        )

    base_name = os.path.splitext(
        os.path.basename(input_file)
    )[0]

    checkpoint_file = os.path.join(
        current_dir,
        f"{base_name}_checkpoint.txt"
    )

    not_imported_file = os.path.join(
        current_dir,
        f"{base_name}_not_imported.xlsx"
    )

    print("=" * 70)
    print("IMPORT ITEM PRICE - DIRECT SQL BATCH MODE")
    print("=" * 70)
    print(f"File       : {input_file}")
    print(f"Batch size : {batch_size}")
    print(f"Checkpoint : {checkpoint_file}")
    print("=" * 70)

    # ==========================================================
    # 2. LOAD EXCEL
    # ==========================================================

    wb = openpyxl.load_workbook(
        input_file,
        read_only=True,
        data_only=True,
    )

    ws = wb.active

    if ws.max_row < 2:
        wb.close()

        frappe.throw(
            "Excel tidak memiliki data."
        )

    total_rows = ws.max_row - 1

    # ==========================================================
    # 3. BACA HEADER
    # ==========================================================

    headers = {}

    for col in range(
        1,
        ws.max_column + 1
    ):

        value = ws.cell(
            row=1,
            column=col
        ).value

        if value:

            header = str(
                value
            ).strip().lower()

            headers[header] = col

    # ==========================================================
    # 4. HELPER CARI KOLOM
    # ==========================================================

    def find_column(*names):

        for name in names:

            name = str(
                name
            ).strip().lower()

            if name in headers:
                return headers[name]

        return None

    # ==========================================================
    # 5. CARI KOLOM
    # ==========================================================

    item_code_col = find_column(
        "item_code",
        "item code",
        "item"
    )

    price_list_col = find_column(
        "price_list",
        "price list"
    )

    price_col = find_column(
        "price_list_rate",
        "price list rate",
        "rate",
        "price"
    )

    currency_col = find_column(
        "currency"
    )

    buying_col = find_column(
        "buying"
    )

    selling_col = find_column(
        "selling"
    )

    valid_from_col = find_column(
        "valid_from",
        "valid from"
    )

    valid_upto_col = find_column(
        "valid_upto",
        "valid upto"
    )

    # ==========================================================
    # 6. VALIDASI KOLOM WAJIB
    # ==========================================================

    if not item_code_col:
        wb.close()

        frappe.throw(
            "Kolom Item Code tidak ditemukan."
        )

    if not price_list_col:
        wb.close()

        frappe.throw(
            "Kolom Price List tidak ditemukan."
        )

    if not price_col:
        wb.close()

        frappe.throw(
            "Kolom Price List Rate / Rate / Price tidak ditemukan."
        )

    # ==========================================================
    # 7. CHECKPOINT
    # ==========================================================

    last_committed_row = 1

    if os.path.exists(checkpoint_file):

        try:

            with open(
                checkpoint_file,
                "r"
            ) as f:

                value = f.read().strip()

                if value:
                    last_committed_row = int(
                        value
                    )

        except Exception as e:

            print(
                f"[WARNING] Checkpoint tidak bisa dibaca: {e}"
            )

            last_committed_row = 1

    start_row = last_committed_row + 1

    if start_row > ws.max_row:

        wb.close()

        print(
            "\nSemua row sudah pernah di-commit."
        )

        return {
            "status": "already_completed",
            "total_rows": total_rows,
            "imported": 0,
            "updated_existing": 0,
            "skipped": 0,
            "last_committed_row": last_committed_row,
            "checkpoint_file": checkpoint_file,
            "not_imported_file": (
                not_imported_file
                if os.path.exists(not_imported_file)
                else None
            ),
        }

    print(
        f"\nResume dari row Excel: {start_row}"
    )

    print(
        f"Row terakhir committed sebelumnya: "
        f"{last_committed_row}"
    )

    # ==========================================================
    # 8. COUNTER
    # ==========================================================

    imported = 0
    updated_existing = 0
    skipped = 0

    total_processed_this_run = 0

    # ==========================================================
    # 9. FAILED EXCEL
    # ==========================================================

    original_headers = [
        ws.cell(
            row=1,
            column=col
        ).value
        for col in range(
            1,
            ws.max_column + 1
        )
    ]

    def save_failed_excel(failed_rows):

        if not failed_rows:
            return

        existing_failed = []

        if os.path.exists(
            not_imported_file
        ):

            try:

                old_wb = openpyxl.load_workbook(
                    not_imported_file
                )

                old_ws = old_wb.active

                for row in old_ws.iter_rows(
                    min_row=2,
                    values_only=True
                ):

                    existing_failed.append(
                        list(row)
                    )

                old_wb.close()

            except Exception as e:

                print(
                    f"[WARNING] Gagal membaca failed Excel lama: {e}"
                )

        failed_wb = openpyxl.Workbook()

        failed_ws = failed_wb.active

        failed_ws.title = "Not Imported"

        failed_ws.append([
            "Original Row",
            "Reason",
            "UOM",
            *original_headers
        ])

        for row in existing_failed:

            failed_ws.append(row)

        for failed in failed_rows:

            failed_ws.append([
                failed["row_number"],
                failed["reason"],
                failed["uom"],
                *failed["data"]
            ])

        failed_wb.save(
            not_imported_file
        )

        failed_wb.close()

    # ==========================================================
    # 10. HELPER VALUE
    # ==========================================================

    def normalize_string(value):

        if value is None:
            return ""

        return str(value).strip()

    def normalize_bool(value):

        if value is None:
            return 0

        if isinstance(value, bool):
            return int(value)

        return (
            1
            if str(value).strip().lower()
            in (
                "1",
                "yes",
                "true",
                "y"
            )
            else 0
        )

    def normalize_price(value):

        if value is None:
            return None

        if isinstance(value, (int, float)):

            return frappe.utils.flt(
                value
            )

        value = str(
            value
        ).strip()

        if not value:
            return None

        value = value.replace(
            "Rp",
            ""
        )

        value = value.replace(
            "rp",
            ""
        )

        value = value.strip()

        if "," in value and "." in value:

            if value.rfind(".") > value.rfind(","):

                value = value.replace(
                    ",",
                    ""
                )

            else:

                value = value.replace(
                    ".",
                    ""
                )

                value = value.replace(
                    ",",
                    "."
                )

        elif "," in value:

            parts = value.split(",")

            if len(parts) == 2 and len(parts[1]) <= 2:

                value = value.replace(
                    ",",
                    "."
                )

            else:

                value = value.replace(
                    ",",
                    ""
                )

        return frappe.utils.flt(
            value
        )

    # ==========================================================
    # 11. LOOP BATCH
    # ==========================================================

    current_row = start_row

    while current_row <= ws.max_row:

        batch_start = current_row

        batch_end = min(
            current_row + batch_size - 1,
            ws.max_row
        )

        print("\n")
        print("=" * 70)
        print(
            f"PROCESS BATCH "
            f"ROW {batch_start} - {batch_end}"
        )
        print("=" * 70)

        batch_imported = 0
        batch_updated = 0
        batch_skipped = 0

        batch_failed_rows = []

        # ======================================================
        # 12. BACA ROW EXCEL KE MEMORY
        # ======================================================

        batch_rows = []

        print(
            f"[READ EXCEL] Mulai membaca row "
            f"{batch_start} - {batch_end}"
        )

        for row_number, row_data in enumerate(
            ws.iter_rows(
                min_row=batch_start,
                max_row=batch_end,
                values_only=True,
            ),
            start=batch_start,
        ):

            batch_rows.append({
                "row_number": row_number,
                "row_data": list(row_data),
            })

            if row_number % 100 == 0:

                print(
                    f"[READ EXCEL] Sudah membaca row {row_number}"
                )

        print(
            f"[READ EXCEL] Selesai membaca "
            f"{len(batch_rows)} row"
        )

        # ======================================================
        # 13. KUMPULKAN ITEM CODE UNTUK 1 QUERY
        # ======================================================

        item_codes = []

        for row in batch_rows:

            row_data = row["row_data"]

            item_code = normalize_string(
                row_data[item_code_col - 1]
            )

            price_list = normalize_string(
                row_data[price_list_col - 1]
            )

            price = row_data[
                price_col - 1
            ]

            if (
                not item_code
                and not price_list
                and (
                    price is None
                    or price == ""
                )
            ):
                continue

            if item_code:

                item_codes.append(
                    item_code
                )

        item_codes = list(
            dict.fromkeys(
                item_codes
            )
        )

        # ======================================================
        # 14. QUERY ITEM SEKALI PER BATCH
        # ======================================================

        existing_items = {}

        if item_codes:

            placeholders = ", ".join(
                ["%s"] * len(item_codes)
            )

            rows = frappe.db.sql(
                f"""
                SELECT
                    name,
                    item_name,
                    stock_uom
                FROM `tabItem`
                WHERE name IN ({placeholders})
                """,
                tuple(item_codes),
                as_dict=True,
            )

            existing_items = {
                row["name"]: {
                    "item_name": row["item_name"],
                    "stock_uom": row["stock_uom"],
                }
                for row in rows
            }

        print(
            f"[ITEM CHECK] "
            f"{len(item_codes)} Item Code "
            f"→ {len(existing_items)} ditemukan"
        )

        # ======================================================
        # 15. VALIDASI EXCEL ROW
        # ======================================================

        valid_rows = []

        for row in batch_rows:

            row_number = row["row_number"]
            row_data = row["row_data"]

            current_uom = ""

            try:

                item_code = normalize_string(
                    row_data[item_code_col - 1]
                )

                price_list = normalize_string(
                    row_data[price_list_col - 1]
                )

                price = row_data[
                    price_col - 1
                ]

                # --------------------------------------------------
                # ROW KOSONG
                # --------------------------------------------------

                if (
                    not item_code
                    and not price_list
                    and (
                        price is None
                        or price == ""
                    )
                ):
                    continue

                # --------------------------------------------------
                # ITEM CODE
                # --------------------------------------------------

                if not item_code:

                    raise Exception(
                        "Item Code kosong"
                    )

                # --------------------------------------------------
                # ITEM EXISTENCE
                # --------------------------------------------------

                if item_code not in existing_items:

                    raise Exception(
                        f"Item '{item_code}' "
                        f"tidak ditemukan"
                    )

                # --------------------------------------------------
                # AMBIL DATA ITEM
                # --------------------------------------------------

                item_info = existing_items[
                    item_code
                ]

                item_name = normalize_string(
                    item_info.get(
                        "item_name"
                    )
                )

                current_uom = normalize_string(
                    item_info.get(
                        "stock_uom"
                    )
                )

                if not current_uom:

                    raise Exception(
                        f"Item '{item_code}' "
                        f"tidak memiliki Stock UOM"
                    )

                # --------------------------------------------------
                # PRICE LIST
                # --------------------------------------------------

                if not price_list:

                    raise Exception(
                        "Price List kosong"
                    )

                # --------------------------------------------------
                # PRICE
                # --------------------------------------------------

                if price is None or price == "":

                    raise Exception(
                        "Rate kosong"
                    )

                price = normalize_price(
                    price
                )

                if price is None:

                    raise Exception(
                        "Rate tidak valid"
                    )

                # --------------------------------------------------
                # CURRENCY
                # --------------------------------------------------

                currency = None

                if currency_col:

                    value = row_data[
                        currency_col - 1
                    ]

                    if value:

                        currency = normalize_string(
                            value
                        )

                # --------------------------------------------------
                # BUYING
                # --------------------------------------------------

                buying = 0

                if buying_col:

                    buying = normalize_bool(
                        row_data[
                            buying_col - 1
                        ]
                    )

                # --------------------------------------------------
                # SELLING
                # --------------------------------------------------

                selling = 0

                if selling_col:

                    selling = normalize_bool(
                        row_data[
                            selling_col - 1
                        ]
                    )

                # --------------------------------------------------
                # VALID FROM
                # --------------------------------------------------

                valid_from = None

                if valid_from_col:

                    value = row_data[
                        valid_from_col - 1
                    ]

                    if value:

                        valid_from = normalize_date(
                            value
                        )

                # --------------------------------------------------
                # VALID UPTO
                # --------------------------------------------------

                valid_upto = None

                if valid_upto_col:

                    value = row_data[
                        valid_upto_col - 1
                    ]

                    if value:

                        valid_upto = normalize_date(
                            value
                        )

                # --------------------------------------------------
                # SIMPAN ROW VALID
                # --------------------------------------------------

                valid_rows.append({
                    "row_number": row_number,
                    "item_code": item_code,
                    "item_name": item_name,
                    "uom": current_uom,
                    "price_list": price_list,
                    "price": price,
                    "currency": currency,
                    "buying": buying,
                    "selling": selling,
                    "valid_from": valid_from,
                    "valid_upto": valid_upto,
                    "data": row_data,
                })

            except Exception as e:

                batch_skipped += 1
                skipped += 1

                batch_failed_rows.append({
                    "row_number": row_number,
                    "reason": str(e),
                    "uom": current_uom,
                    "data": row_data,
                })

                print(
                    f"[FAILED] Row {row_number}: "
                    f"{str(e)}"
                )

                continue

        # ======================================================
        # 16. CEK DUPLIKAT DI DALAM EXCEL
        # ======================================================

        unique_rows = {}

        for row in valid_rows:

            key = (
                row["item_code"],
                row["price_list"],
                row["uom"],
            )

            unique_rows[key] = row

        valid_rows = list(
            unique_rows.values()
        )

        print(
            f"[VALID] "
            f"{len(valid_rows)} row siap diproses"
        )

        # ======================================================
        # 17. QUERY EXISTING ITEM PRICE SEKALI PER BATCH
        # ======================================================

        existing_prices = {}

        if valid_rows:

            item_price_conditions = []
            item_price_values = []

            for row in valid_rows:

                item_price_conditions.append(
                    """
                    (
                        item_code = %s
                        AND price_list = %s
                        AND uom = %s
                    )
                    """
                )

                item_price_values.extend([
                    row["item_code"],
                    row["price_list"],
                    row["uom"],
                ])

            where_clause = " OR ".join(
                item_price_conditions
            )

            existing_price_rows = frappe.db.sql(
                f"""
                SELECT
                    name,
                    item_code,
                    price_list,
                    uom
                FROM `tabItem Price`
                WHERE {where_clause}
                """,
                tuple(item_price_values),
                as_dict=True,
            )

            for existing in existing_price_rows:

                key = (
                    existing["item_code"],
                    existing["price_list"],
                    existing["uom"],
                )

                existing_prices[key] = (
                    existing["name"]
                )

        print(
            f"[ITEM PRICE CHECK] "
            f"{len(existing_prices)} existing"
        )

        # ======================================================
        # 18. INSERT / UPDATE DIRECT SQL
        # ======================================================

        try:

            for row in valid_rows:

                key = (
                    row["item_code"],
                    row["price_list"],
                    row["uom"],
                )

                existing_name = existing_prices.get(
                    key
                )

                # ==================================================
                # UPDATE
                # ==================================================

                if existing_name:

                    frappe.db.sql(
                        """
                        UPDATE `tabItem Price`
                        SET
                            item_name = %s,
                            price_list_rate = %s,
                            currency = %s,
                            buying = %s,
                            selling = %s,
                            valid_from = %s,
                            valid_upto = %s,
                            modified = NOW(),
                            modified_by = %s
                        WHERE name = %s
                        """,
                        (
                            row["item_name"],
                            row["price"],
                            row["currency"],
                            row["buying"],
                            row["selling"],
                            row["valid_from"],
                            row["valid_upto"],
                            frappe.session.user,
                            existing_name,
                        )
                    )

                    batch_updated += 1
                    updated_existing += 1

                    batch_imported += 1
                    imported += 1

                    print(
                        f"[UPDATE] Row {row['row_number']}: "
                        f"{row['item_code']} | "
                        f"{row['price_list']} | "
                        f"{row['uom']} | "
                        f"{row['price']}"
                    )

                # ==================================================
                # INSERT
                # ==================================================

                else:

                    item_price_name = frappe.generate_hash(
                        length=10
                    )

                    frappe.db.sql(
                        """
                        INSERT INTO `tabItem Price`
                        (
                            name,
                            creation,
                            modified,
                            modified_by,
                            owner,
                            docstatus,
                            idx,
                            item_code,
                            item_name,
                            price_list,
                            price_list_rate,
                            currency,
                            uom,
                            buying,
                            selling,
                            valid_from,
                            valid_upto
                        )
                        VALUES
                        (
                            %s,
                            NOW(),
                            NOW(),
                            %s,
                            %s,
                            0,
                            0,
                            %s,
                            %s,
                            %s,
                            %s,
                            %s,
                            %s,
                            %s,
                            %s,
                            %s,
                            %s
                        )
                        """,
                        (
                            item_price_name,
                            frappe.session.user,
                            frappe.session.user,
                            row["item_code"],
                            row["item_name"],
                            row["price_list"],
                            row["price"],
                            row["currency"],
                            row["uom"],
                            row["buying"],
                            row["selling"],
                            row["valid_from"],
                            row["valid_upto"],
                        )
                    )

                    batch_imported += 1
                    imported += 1

                    print(
                        f"[INSERT] Row {row['row_number']}: "
                        f"{row['item_code']} | "
                        f"{row['price_list']} | "
                        f"{row['uom']} | "
                        f"{row['price']}"
                    )

        except Exception as e:

            frappe.db.rollback()

            print("\n" + "!" * 70)
            print("[DATABASE ERROR]")
            print(str(e))
            print(
                f"Batch row {batch_start} - {batch_end} "
                f"di-ROLLBACK."
            )
            print(
                "Checkpoint TIDAK dipindahkan."
            )
            print("!" * 70)

            wb.close()

            raise

        # ======================================================
        # 19. SIMPAN FAILED EXCEL
        # ======================================================

        if batch_failed_rows:

            try:

                save_failed_excel(
                    batch_failed_rows
                )

                print(
                    f"[FAILED EXCEL] "
                    f"{len(batch_failed_rows)} row dicatat."
                )

            except Exception as e:

                print(
                    f"[WARNING] "
                    f"Gagal membuat Excel failed: "
                    f"{str(e)}"
                )

        # ======================================================
        # 20. COMMIT BATCH
        # ======================================================

        if commit:

            try:

                frappe.db.commit()

                print("\n" + "-" * 70)
                print(
                    f"[COMMIT BERHASIL] "
                    f"Row {batch_start} - {batch_end}"
                )
                print(
                    f"Insert/Update : {batch_imported}"
                )
                print(
                    f"Failed        : {batch_skipped}"
                )
                print("-" * 70)

            except Exception as e:

                frappe.db.rollback()

                print("\n" + "!" * 70)
                print("[COMMIT ERROR]")
                print(str(e))
                print(
                    "Batch di-ROLLBACK."
                )
                print(
                    "Checkpoint tidak diubah."
                )
                print("!" * 70)

                wb.close()

                return {
                    "status": "commit_failed",
                    "failed_batch_start": batch_start,
                    "failed_batch_end": batch_end,
                    "total_rows": total_rows,
                    "imported": imported,
                    "updated_existing": updated_existing,
                    "skipped": skipped,
                    "last_committed_row": last_committed_row,
                    "not_imported_file": (
                        not_imported_file
                        if os.path.exists(
                            not_imported_file
                        )
                        else None
                    ),
                }

        # ======================================================
        # 21. CHECKPOINT
        # ======================================================

        try:

            with open(
                checkpoint_file,
                "w"
            ) as f:

                f.write(
                    str(batch_end)
                )

            last_committed_row = batch_end

            print(
                f"[CHECKPOINT] "
                f"Row {batch_end} sudah aman."
            )

        except Exception as e:

            print("\n" + "!" * 70)
            print(
                "[WARNING] Database sudah COMMIT,"
            )
            print(
                "tetapi checkpoint gagal disimpan."
            )
            print(str(e))
            print(
                "Jangan hapus database."
            )
            print(
                "Checkpoint perlu diperbaiki "
                "sebelum menjalankan ulang."
            )
            print("!" * 70)

            wb.close()

            return {
                "status": "checkpoint_failed",
                "last_committed_row": last_committed_row,
                "total_rows": total_rows,
                "imported": imported,
                "updated_existing": updated_existing,
                "skipped": skipped,
                "checkpoint_file": checkpoint_file,
                "not_imported_file": (
                    not_imported_file
                    if os.path.exists(
                        not_imported_file
                    )
                    else None
                ),
            }

        # ======================================================
        # 22. NEXT BATCH
        # ======================================================

        total_processed_this_run += (
            batch_end - batch_start + 1
        )

        current_row = batch_end + 1

        print(
            f"[NEXT] Lanjut ke row {current_row}"
        )

    # ==========================================================
    # 23. SELESAI
    # ==========================================================

    wb.close()

    print("\n")
    print("=" * 70)
    print("IMPORT ITEM PRICE SELESAI")
    print("=" * 70)

    print(
        f"Total row Excel        : {total_rows}"
    )

    print(
        f"Row diproses kali ini  : "
        f"{total_processed_this_run}"
    )

    print(
        f"Berhasil dibuat/update : {imported}"
    )

    print(
        f"Existing di-update     : {updated_existing}"
    )

    print(
        f"Gagal / tidak import   : {skipped}"
    )

    print(
        f"Last committed row     : "
        f"{last_committed_row}"
    )

    if os.path.exists(
        not_imported_file
    ):

        print(
            "\nFile daftar gagal:"
        )

        print(
            not_imported_file
        )

    else:

        print(
            "\nSemua data berhasil diproses."
        )

    print("=" * 70)

    return {
        "status": "success",
        "total_rows": total_rows,
        "processed_this_run": total_processed_this_run,
        "imported": imported,
        "updated_existing": updated_existing,
        "skipped": skipped,
        "last_committed_row": last_committed_row,
        "checkpoint_file": checkpoint_file,
        "not_imported_file": (
            not_imported_file
            if os.path.exists(
                not_imported_file
            )
            else None
        ),
    }

def update_item_pricelist(
    filename="Item Price (Update) 5-10-2026.xlsx",
    batch_size=1000
):
    import os
    from datetime import datetime, date

    import openpyxl
    import frappe

    # ============================================================
    # PATH
    # ============================================================
    file_path = os.path.join(
        os.path.dirname(__file__),
        filename
    )

    checkpoint_file = os.path.join(
        os.path.dirname(__file__),
        "item_price_update_checkpoint.txt"
    )

    failed_file = os.path.join(
        os.path.dirname(__file__),
        "Item Price Update Failed.xlsx"
    )

    if not os.path.exists(file_path):
        frappe.throw(
            f"File tidak ditemukan: {file_path}"
        )

    # ============================================================
    # CHECKPOINT
    # ============================================================
    start_row = 2

    if os.path.exists(checkpoint_file):
        try:
            with open(checkpoint_file, "r") as f:
                checkpoint = f.read().strip()

            if checkpoint:
                start_row = int(checkpoint)

        except Exception:
            start_row = 2

    # ============================================================
    # LOAD EXCEL
    # ============================================================
    wb = openpyxl.load_workbook(
        file_path,
        read_only=True,
        data_only=True
    )

    ws = wb.active

    # ============================================================
    # HEADER
    # ============================================================
    header_row = next(
        ws.iter_rows(
            min_row=1,
            max_row=1,
            values_only=True
        )
    )

    headers = {}

    for col_idx, value in enumerate(
        header_row,
        start=1
    ):
        if value is not None:
            headers[
                str(value).strip().lower()
            ] = col_idx

    def get_col(*names):
        for name in names:
            col = headers.get(
                name.lower()
            )

            if col:
                return col

        return None

    # ============================================================
    # GET COLUMNS
    # ============================================================
    id_col = get_col(
        "id",
        "name"
    )

    item_code_col = get_col(
        "item code",
        "item_code"
    )

    item_name_col = get_col(
        "item name",
        "item_name"
    )

    uom_col = get_col(
        "uom"
    )

    price_list_col = get_col(
        "price list",
        "price_list"
    )

    buying_col = get_col(
        "buying"
    )

    selling_col = get_col(
        "selling"
    )

    valid_from_col = get_col(
        "valid from",
        "valid_from"
    )

    valid_upto_col = get_col(
        "valid upto",
        "valid_upto"
    )

    # ============================================================
    # REQUIRED COLUMNS
    # ============================================================
    required_columns = {
        "ID": id_col,
        "Item Code": item_code_col,
        "UOM": uom_col,
        "Price List": price_list_col,
        "Buying": buying_col,
        "Selling": selling_col,
        "Valid From": valid_from_col,
    }

    missing_columns = [
        name
        for name, col in required_columns.items()
        if col is None
    ]

    if missing_columns:
        frappe.throw(
            "Kolom Excel berikut tidak ditemukan: "
            + ", ".join(missing_columns)
        )

    # ============================================================
    # NORMALIZE BOOLEAN
    # ============================================================
    def normalize_bool(value):

        if value is None:
            return 0

        if isinstance(value, bool):
            return 1 if value else 0

        if isinstance(value, (int, float)):
            return 1 if value != 0 else 0

        value = str(value).strip().lower()

        if value in (
            "1",
            "yes",
            "true",
            "y",
            "ya",
        ):
            return 1

        return 0

    # ============================================================
    # NORMALIZE DATE
    # ============================================================
    def normalize_date(value):

        if value is None:
            return None

        if isinstance(value, datetime):
            return value.date()

        if isinstance(value, date):
            return value

        if isinstance(value, str):

            value = value.strip()

            if not value:
                return None

            date_formats = [
                "%d-%m-%Y",
                "%d/%m/%Y",
                "%Y-%m-%d",
                "%Y/%m/%d",
            ]

            for fmt in date_formats:
                try:
                    return datetime.strptime(
                        value,
                        fmt
                    ).date()

                except ValueError:
                    continue

        raise ValueError(
            f"Format Valid From tidak valid: {value}"
        )

    # ============================================================
    # LOAD UOM
    # ============================================================
    uom_list = frappe.get_all(
        "UOM",
        pluck="name"
    )

    uom_map = {
        str(uom).strip().lower(): uom
        for uom in uom_list
    }

    # ============================================================
    # FAILED ROWS
    # ============================================================
    failed_rows = []

    # ============================================================
    # TOTAL ROWS
    # ============================================================
    total_rows = ws.max_row - 1

    print("=" * 100)
    print("UPDATE ITEM PRICE")
    print("=" * 100)
    print(f"File       : {file_path}")
    print(f"Start row  : {start_row}")
    print(f"Total rows : {total_rows}")
    print(f"Batch size : {batch_size}")
    print("=" * 100)

    # ============================================================
    # PROCESS BATCH
    # ============================================================
    current_row = start_row

    while current_row <= ws.max_row:

        batch_end = min(
            current_row + batch_size - 1,
            ws.max_row
        )

        print(
            f"\nProcessing rows "
            f"{current_row} - {batch_end}"
        )

        # ========================================================
        # READ BATCH
        # ========================================================
        batch_rows = []

        for row in ws.iter_rows(
            min_row=current_row,
            max_row=batch_end,
            values_only=True
        ):
            batch_rows.append(row)

        # ========================================================
        # COLLECT ITEM PRICE IDs
        # ========================================================
        item_price_ids = []

        for row in batch_rows:

            item_price_id = row[
                id_col - 1
            ]

            if item_price_id is None:
                continue

            item_price_id = str(
                item_price_id
            ).strip()

            if item_price_id:
                item_price_ids.append(
                    item_price_id
                )

        item_price_ids = list(
            dict.fromkeys(
                item_price_ids
            )
        )

        # ========================================================
        # GET ITEM PRICE BY ID
        # ========================================================
        existing_prices = {}

        if item_price_ids:

            placeholders = ", ".join(
                ["%s"] * len(item_price_ids)
            )

            query = f"""
                SELECT
                    name,
                    item_code,
                    item_name,
                    price_list,
                    uom,
                    price_list_rate,
                    buying,
                    selling,
                    valid_from,
                    valid_upto
                FROM `tabItem Price`
                WHERE name IN ({placeholders})
            """

            result = frappe.db.sql(
                query,
                tuple(item_price_ids),
                as_dict=True
            )

            existing_prices = {
                row["name"]: row
                for row in result
            }

        # ========================================================
        # UPDATE BATCH
        # ========================================================
        batch_updated = 0
        batch_failed = 0

        try:

            for excel_row_number, row in zip(
                range(
                    current_row,
                    batch_end + 1
                ),
                batch_rows
            ):

                # =================================================
                # GET EXCEL DATA
                # =================================================
                item_price_id = row[
                    id_col - 1
                ]

                item_code = row[
                    item_code_col - 1
                ]

                item_name = (
                    row[item_name_col - 1]
                    if item_name_col
                    else None
                )

                excel_uom = row[
                    uom_col - 1
                ]

                price_list = row[
                    price_list_col - 1
                ]

                buying = normalize_bool(
                    row[buying_col - 1]
                )

                selling = normalize_bool(
                    row[selling_col - 1]
                )

                valid_from_raw = row[
                    valid_from_col - 1
                ]

                # =================================================
                # VALIDATE ID
                # =================================================
                if item_price_id is None:

                    failed_row = list(row)

                    failed_row.append(
                        "ID kosong"
                    )

                    failed_rows.append(
                        failed_row
                    )

                    batch_failed += 1

                    continue

                item_price_id = str(
                    item_price_id
                ).strip()

                if not item_price_id:

                    failed_row = list(row)

                    failed_row.append(
                        "ID kosong"
                    )

                    failed_rows.append(
                        failed_row
                    )

                    batch_failed += 1

                    continue

                # =================================================
                # FIND ITEM PRICE
                # =================================================
                db_price = existing_prices.get(
                    item_price_id
                )

                if not db_price:

                    failed_row = list(row)

                    failed_row.append(
                        f"Item Price ID tidak ditemukan: "
                        f"{item_price_id}"
                    )

                    failed_rows.append(
                        failed_row
                    )

                    batch_failed += 1

                    print(
                        f"[FAILED] Row "
                        f"{excel_row_number}: "
                        f"ID tidak ditemukan "
                        f"{item_price_id}"
                    )

                    continue

                # =================================================
                # VALIDATE ITEM CODE
                # =================================================
                if item_code is None:

                    failed_row = list(row)

                    failed_row.append(
                        "Item Code kosong"
                    )

                    failed_rows.append(
                        failed_row
                    )

                    batch_failed += 1

                    continue

                item_code = str(
                    item_code
                ).strip()

                db_item_code = str(
                    db_price["item_code"]
                ).strip()

                if item_code != db_item_code:

                    failed_row = list(row)

                    failed_row.append(
                        "Item Code Excel tidak sama "
                        f"dengan DB. "
                        f"Excel={item_code}, "
                        f"DB={db_item_code}"
                    )

                    failed_rows.append(
                        failed_row
                    )

                    batch_failed += 1

                    print(
                        f"[FAILED] Row "
                        f"{excel_row_number}: "
                        f"Item Code mismatch "
                        f"ID={item_price_id}"
                    )

                    continue

                # =================================================
                # VALIDATE PRICE LIST
                # =================================================
                if price_list is None:

                    failed_row = list(row)

                    failed_row.append(
                        "Price List kosong"
                    )

                    failed_rows.append(
                        failed_row
                    )

                    batch_failed += 1

                    continue

                price_list = str(
                    price_list
                ).strip()

                db_price_list = str(
                    db_price["price_list"]
                ).strip()

                if price_list != db_price_list:

                    failed_row = list(row)

                    failed_row.append(
                        "Price List Excel tidak sama "
                        f"dengan DB. "
                        f"Excel={price_list}, "
                        f"DB={db_price_list}"
                    )

                    failed_rows.append(
                        failed_row
                    )

                    batch_failed += 1

                    print(
                        f"[FAILED] Row "
                        f"{excel_row_number}: "
                        f"Price List mismatch "
                        f"ID={item_price_id}"
                    )

                    continue

                # =================================================
                # VALIDATE UOM
                # =================================================
                if excel_uom is None:

                    failed_row = list(row)

                    failed_row.append(
                        "UOM kosong"
                    )

                    failed_rows.append(
                        failed_row
                    )

                    batch_failed += 1

                    continue

                excel_uom_key = str(
                    excel_uom
                ).strip().lower()

                system_uom = uom_map.get(
                    excel_uom_key
                )

                if not system_uom:

                    failed_row = list(row)

                    failed_row.append(
                        f"UOM tidak ditemukan: "
                        f"{excel_uom}"
                    )

                    failed_rows.append(
                        failed_row
                    )

                    batch_failed += 1

                    print(
                        f"[FAILED] Row "
                        f"{excel_row_number}: "
                        f"UOM tidak ditemukan "
                        f"{excel_uom}"
                    )

                    continue

                db_uom = (
                    str(
                        db_price["uom"]
                    ).strip()
                    if db_price["uom"]
                    else ""
                )

                if system_uom != db_uom:

                    failed_row = list(row)

                    failed_row.append(
                        "UOM Excel tidak sama "
                        f"dengan DB. "
                        f"Excel={system_uom}, "
                        f"DB={db_uom}"
                    )

                    failed_rows.append(
                        failed_row
                    )

                    batch_failed += 1

                    print(
                        f"[FAILED] Row "
                        f"{excel_row_number}: "
                        f"UOM mismatch "
                        f"ID={item_price_id}"
                    )

                    continue

                # =================================================
                # VALIDATE VALID FROM
                # =================================================
                try:

                    valid_from = normalize_date(
                        valid_from_raw
                    )

                except Exception as e:

                    failed_row = list(row)

                    failed_row.append(
                        str(e)
                    )

                    failed_rows.append(
                        failed_row
                    )

                    batch_failed += 1

                    print(
                        f"[FAILED] Row "
                        f"{excel_row_number}: "
                        f"{str(e)}"
                    )

                    continue

                if not valid_from:

                    failed_row = list(row)

                    failed_row.append(
                        "Valid From kosong"
                    )

                    failed_rows.append(
                        failed_row
                    )

                    batch_failed += 1

                    continue

                # =================================================
                # UPDATE ITEM PRICE
                # =================================================
                frappe.db.sql(
                    """
                    UPDATE `tabItem Price`
                    SET
                        buying = %s,
                        selling = %s,
                        valid_from = %s,
                        modified = NOW(),
                        modified_by = %s
                    WHERE name = %s
                    """,
                    (
                        buying,
                        selling,
                        valid_from,
                        frappe.session.user,
                        item_price_id
                    )
                )

                batch_updated += 1

                print(
                    f"[UPDATED] Row "
                    f"{excel_row_number}: "
                    f"{item_price_id} | "
                    f"{item_code} | "
                    f"{price_list} | "
                    f"{system_uom} | "
                    f"buying={buying} | "
                    f"selling={selling} | "
                    f"valid_from={valid_from}"
                )

            # ====================================================
            # SAVE FAILED EXCEL
            # ====================================================
            if batch_failed > 0:

                failed_wb = openpyxl.Workbook()
                failed_ws = failed_wb.active

                header_values = list(
                    header_row
                )

                header_values.append(
                    "Error"
                )

                failed_ws.append(
                    header_values
                )

                for failed_row in failed_rows:
                    failed_ws.append(
                        failed_row
                    )

                failed_wb.save(
                    failed_file
                )

                failed_wb.close()

            # ====================================================
            # COMMIT
            # ====================================================
            frappe.db.commit()

            # ====================================================
            # CHECKPOINT
            # ====================================================
            next_row = batch_end + 1

            with open(
                checkpoint_file,
                "w"
            ) as f:
                f.write(
                    str(next_row)
                )

            print("-" * 100)
            print(
                f"BATCH SELESAI | "
                f"Updated={batch_updated} | "
                f"Failed={batch_failed}"
            )
            print(
                f"Checkpoint -> row {next_row}"
            )
            print("-" * 100)

        except Exception as e:

            frappe.db.rollback()

            frappe.log_error(
                frappe.get_traceback(),
                "Update Item Price Failed"
            )

            print(
                f"[ERROR] Batch "
                f"{current_row}-{batch_end}: "
                f"{str(e)}"
            )

            raise

        # ========================================================
        # NEXT BATCH
        # ========================================================
        current_row = batch_end + 1

    # ============================================================
    # CLOSE EXCEL
    # ============================================================
    wb.close()

    # ============================================================
    # SUMMARY
    # ============================================================
    print("\n")
    print("=" * 100)
    print("UPDATE ITEM PRICE SELESAI")
    print("=" * 100)

    if failed_rows:
        print(
            f"Total failed rows : {len(failed_rows)}"
        )
        print(
            f"Failed file       : {failed_file}"
        )
    else:
        print(
            "Tidak ada row yang gagal."
        )

    print("=" * 100)

    return {
        "status": "success",
        "total_rows": total_rows,
        "failed_rows": len(failed_rows),
        "failed_file": (
            failed_file
            if failed_rows
            else None
        ),
        "checkpoint_file": checkpoint_file,
    }

# ==============================================================
# NORMALIZE DATE
# ==============================================================

def normalize_date(value):
    """
    Normalisasi tanggal Excel menjadi YYYY-MM-DD.
    """

    if value is None:
        return None

    if isinstance(value, datetime):
        return value.strftime(
            "%Y-%m-%d"
        )

    if isinstance(value, date):
        return value.strftime(
            "%Y-%m-%d"
        )

    value = str(value).strip()

    if not value:
        return None

    formats = [
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%d-%m-%y",
        "%d/%m/%y",
    ]

    for fmt in formats:

        try:

            return datetime.strptime(
                value,
                fmt
            ).strftime(
                "%Y-%m-%d"
            )

        except ValueError:
            continue

    raise ValueError(
        f"Format tanggal tidak dikenali: {value}"
    )






def update_item_price_valid_from():

    import frappe
    import os
    import json
    from openpyxl import load_workbook
    from datetime import datetime


    BATCH_SIZE = 500

    EXCEL_FILE = "Item Price (Update tgl ke 5 sept 2026).xlsx"
    CHECKPOINT_FILE = "item_price_checkpoint.json"


    base_dir = os.path.dirname(os.path.abspath(__file__))

    excel_path = os.path.join(base_dir, EXCEL_FILE)
    checkpoint_path = os.path.join(base_dir, CHECKPOINT_FILE)

    if not os.path.exists(excel_path):
        frappe.throw(f"File Excel tidak ditemukan: {excel_path}")

    # ==========================================================
    # LOAD CHECKPOINT
    # ==========================================================

    checkpoint = {
        "last_row": 1,
        "updated": 0,
        "not_found": 0,
        "skipped": 0,
    }

    if os.path.exists(checkpoint_path):
        try:
            with open(checkpoint_path, "r") as f:
                checkpoint.update(json.load(f))

            print(
                f"CHECKPOINT ditemukan. "
                f"Melanjutkan dari row {checkpoint['last_row'] + 1}"
            )

        except Exception as e:
            print(f"WARNING: Gagal membaca checkpoint: {e}")

    start_row = checkpoint["last_row"] + 1

    # ==========================================================
    # LOAD EXCEL
    # ==========================================================

    wb = load_workbook(
        excel_path,
        read_only=True,
        data_only=True
    )

    ws = wb.active

    total_rows = ws.max_row

    print(f"Total row Excel : {total_rows}")
    print(f"Mulai dari row  : {start_row}")
    print(f"Batch size      : {BATCH_SIZE}")

    # ==========================================================
    # PROCESS
    # ==========================================================

    batch_count = 0

    for row_number, row in enumerate(
        ws.iter_rows(min_row=2, values_only=True),
        start=2
    ):
        # Lewati data yang sudah diproses
        if row_number < start_row:
            continue

        item_price_id = row[0]
        valid_from = row[1]

        # ------------------------------------------------------
        # DATA KOSONG
        # ------------------------------------------------------

        if not item_price_id or not valid_from:
            checkpoint["skipped"] += 1
            checkpoint["last_row"] = row_number

            batch_count += 1

        else:

            # --------------------------------------------------
            # CONVERT DATE
            # --------------------------------------------------

            if isinstance(valid_from, datetime):
                valid_from = valid_from.date()

            elif isinstance(valid_from, str):
                try:
                    valid_from = datetime.strptime(
                        valid_from.strip(),
                        "%d-%m-%Y"
                    ).date()

                except ValueError:
                    print(
                        f"SKIP row {row_number}: "
                        f"tanggal invalid "
                        f"{item_price_id} = {valid_from}"
                    )

                    checkpoint["skipped"] += 1
                    checkpoint["last_row"] = row_number

                    batch_count += 1
                    continue

            # --------------------------------------------------
            # CEK ITEM PRICE
            # --------------------------------------------------

            if not frappe.db.exists(
                "Item Price",
                item_price_id
            ):
                print(
                    f"NOT FOUND row {row_number}: "
                    f"{item_price_id}"
                )

                checkpoint["not_found"] += 1

            else:

                # --------------------------------------------------
                # UPDATE
                # --------------------------------------------------

                frappe.db.set_value(
                    "Item Price",
                    item_price_id,
                    "valid_from",
                    valid_from,
                    update_modified=False
                )

                checkpoint["updated"] += 1

                print(
                    f"[{row_number}/{total_rows}] "
                    f"UPDATED: {item_price_id} "
                    f"-> {valid_from}"
                )

            checkpoint["last_row"] = row_number
            batch_count += 1

        # ======================================================
        # CHECKPOINT
        # ======================================================

        if batch_count >= BATCH_SIZE:

            # Commit database
            frappe.db.commit()

            # Simpan checkpoint
            with open(checkpoint_path, "w") as f:
                json.dump(
                    checkpoint,
                    f,
                    indent=4
                )

            print(
                "\n"
                "====================================\n"
                f"CHECKPOINT SAVED\n"
                f"Last row : {checkpoint['last_row']}\n"
                f"Updated  : {checkpoint['updated']}\n"
                f"Not Found: {checkpoint['not_found']}\n"
                f"Skipped  : {checkpoint['skipped']}\n"
                "====================================\n"
            )

            batch_count = 0

    # ==========================================================
    # FINAL COMMIT
    # ==========================================================

    frappe.db.commit()

    checkpoint["last_row"] = total_rows

    with open(checkpoint_path, "w") as f:
        json.dump(
            checkpoint,
            f,
            indent=4
        )

    wb.close()

    print("\n====================================")
    print("SELESAI")
    print(f"Total row : {total_rows}")
    print(f"Updated   : {checkpoint['updated']}")
    print(f"Not Found : {checkpoint['not_found']}")
    print(f"Skipped   : {checkpoint['skipped']}")
    print(f"Checkpoint: {checkpoint_path}")
    print("====================================")



def test_item():

    item_code = "0000063512182"

    print("\n=== ITEM ===")
    print(frappe.db.get_value(
        "Item",
        item_code,
        ["name", "item_code", "item_name", "disabled"],
        as_dict=True
    ))

    print("\n=== BIN ===")
    bins = frappe.db.sql("""
        SELECT
            b.name,
            b.item_code,
            b.warehouse,
            b.actual_qty
        FROM `tabBin` b
        WHERE b.item_code = %s
    """, item_code, as_dict=True)

    for row in bins:
        print(row)

    print("\n=== WAREHOUSE ===")
    for row in bins:
        print(frappe.db.get_value(
            "Warehouse",
            row.warehouse,
            ["name", "warehouse_name", "company", "disabled"],
            as_dict=True
        ))

    print("\n=== SEARCH BJB DIRECT ===")

    result = frappe.db.sql("""
        SELECT
            i.name,
            i.item_code,
            i.item_name,
            w.name AS warehouse,
            w.company,
            b.actual_qty
        FROM `tabItem` i
        INNER JOIN `tabBin` b
            ON b.item_code = i.item_code
        INNER JOIN `tabWarehouse` w
            ON w.name = b.warehouse
        WHERE i.item_code = %s
          AND i.disabled = 0
          AND w.company = 'BJB'
    """, item_code, as_dict=True)

    print("JUMLAH:", len(result))

    for row in result:
        print(row)


def test_pi_series_max():
    series = "PI-AEP2609"

    print("\n=== SERIES ===")
    print(series)

    print("\n=== TAB SERIES ===")

    data = frappe.db.sql("""
        SELECT name, current
        FROM `tabSeries`
        WHERE name = %s
    """, (series,), as_dict=True)

    print(data)

    print("\n=== MAX NOMOR PI ===")

    result = frappe.db.sql("""
        SELECT
            MAX(
                CAST(
                    SUBSTRING(name, %s)
                    AS UNSIGNED
                )
            ) AS max_number
        FROM `tabPurchase Invoice`
        WHERE name LIKE %s
    """, (len(series) + 1, series + "%"), as_dict=True)

    print(result)

    print("\n=== PI TERBESAR ===")

    result = frappe.db.sql("""
        SELECT name, company, creation
        FROM `tabPurchase Invoice`
        WHERE name LIKE %s
        ORDER BY CAST(SUBSTRING(name, %s) AS UNSIGNED) DESC
        LIMIT 10
    """, (series + "%", len(series) + 1), as_dict=True)

    for row in result:
        print(row)


import frappe


def fix_all_series():

	series_config = {
		"Purchase Invoice": "tabPurchase Invoice",
		"Purchase Receipt": "tabPurchase Receipt",
		"Purchase Order": "tabPurchase Order",
		"Sales Invoice": "tabSales Invoice",
		"Delivery Note": "tabDelivery Note",
		"Sales Order": "tabSales Order",
	}

	print("\n========================================")
	print("FIX ALL TAB SERIES")
	print("========================================")

	updated = 0
	skipped = 0
	not_found = 0

	# Ambil SEMUA tabSeries
	all_series = frappe.db.sql("""
		SELECT name, current
		FROM `tabSeries`
		ORDER BY name
	""", as_dict=True)

	print("Total tabSeries:", len(all_series))

	for series_data in all_series:

		series = series_data.name
		current = int(series_data.current or 0)

		print("\n========================================")
		print("SERIES:", series)
		print("CURRENT:", current)
		print("========================================")

		# Tentukan DocType berdasarkan prefix series
		if series.startswith(("PI-", "PR-")):
			doctype = "Purchase Invoice"
			table = "tabPurchase Invoice"

		elif series.startswith("PT-"):
			doctype = "Purchase Receipt"
			table = "tabPurchase Receipt"

		elif series.startswith("PO-"):
			doctype = "Purchase Order"
			table = "tabPurchase Order"

		elif series.startswith(("SI-", "SR-")):
			doctype = "Sales Invoice"
			table = "tabSales Invoice"

		elif series.startswith("DN-"):
			doctype = "Delivery Note"
			table = "tabDelivery Note"

		elif series.startswith("SO-"):
			doctype = "Sales Order"
			table = "tabSales Order"

		else:
			print("SKIP - prefix tidak dikenali")
			skipped += 1
			continue

		print("DOCTYPE:", doctype)

		# tabSeries menggunakan "." sebagai separator.
		# Nama dokumen tidak memiliki ".".
		document_prefix = series.rstrip(".")

		print("Document prefix:", document_prefix)

		result = frappe.db.sql(f"""
			SELECT
				MAX(
					CAST(
						SUBSTRING(name, %s)
						AS UNSIGNED
					)
				) AS max_number
			FROM `{table}`
			WHERE name LIKE %s
		""", (
			len(document_prefix) + 1,
			document_prefix + "%"
		), as_dict=True)

		max_number = int(result[0].max_number or 0)

		print("Max number:", max_number)

		if max_number <= current:
			print("SKIP")
			print("Current sudah >= nomor terbesar.")
			skipped += 1
			continue

		print("\n=== UPDATE ===")
		print("Before:", current)
		print("After :", max_number)

		frappe.db.sql("""
			UPDATE `tabSeries`
			SET current = %s
			WHERE name = %s
		""", (
			max_number,
			series
		))

		updated += 1

	frappe.db.commit()

	print("\n========================================")
	print("SELESAI")
	print("========================================")
	print("Updated:", updated)
	print("Skipped:", skipped)
	print("========================================")

	# Verifikasi
	print("\n=== TAB SERIES SETELAH UPDATE ===")

	result = frappe.db.sql("""
		SELECT name, current
		FROM `tabSeries`
		ORDER BY name
	""", as_dict=True)

	for row in result:
		print(row)

	print("\nSUCCESS")


import re

@frappe.whitelist()
def preview_autoname_series():

	doctypes = {
		"Purchase Invoice": ["PI-", "PR-"],
		"Purchase Receipt": ["PT-"],
		"Purchase Order": ["PO-"],
		"Sales Invoice": ["SI-", "SR-"],
		"Delivery Note": ["DN-"],
		"Sales Order": ["SO-"],
	}

	print("\n========================================")
	print("PREVIEW AUTONAME SERIES")
	print("TIDAK ADA UPDATE / INSERT")
	print("========================================")

	for doctype, prefixes in doctypes.items():

		print(f"\n\n========== {doctype} ==========")

		series_map = {}

		for prefix in prefixes:

			rows = frappe.db.sql(
				f"""
				SELECT name
				FROM `tab{doctype}`
				WHERE name LIKE %s
				""",
				(prefix + "%",),
				as_dict=True
			)

			for row in rows:

				name = row.name

				# Autoname menggunakan #####
				match = re.match(r"^(.*?)(\d{5})$", name)

				if not match:
					continue

				series = match.group(1).rstrip(".")
				number = int(match.group(2))

				if series not in series_map:
					series_map[series] = number
				else:
					series_map[series] = max(
						series_map[series],
						number
					)

		if not series_map:
			print("Tidak ada series ditemukan.")
			continue

		for series in sorted(series_map):

			max_number = series_map[series]

			existing = frappe.db.sql(
				"""
				SELECT name, current
				FROM `tabSeries`
				WHERE name = %s
				""",
				(series,),
				as_dict=True
			)

			if existing:
				current = int(existing[0].current or 0)

				if max_number > current:
					status = "UPDATE"
				else:
					status = "OK"

				print(
					f"{status:6} | "
					f"{series:25} | "
					f"tabSeries={current:<8} | "
					f"max_doc={max_number}"
				)

			else:
				print(
					f"CREATE | "
					f"{series:25} | "
					f"tabSeries={'TIDAK ADA':<8} | "
					f"max_doc={max_number}"
				)

	print("\n========================================")
	print("PREVIEW SELESAI")
	print("Tidak ada perubahan database.")
	print("========================================")


@frappe.whitelist()
def fix_pi_bjm42609():

	series = "PI-BJM42609"

	print("\n========================================")
	print("FIX SERIES")
	print("========================================")
	print("Series :", series)

	# Cari nomor PI terbesar
	result = frappe.db.sql(
		"""
		SELECT MAX(
			CAST(
				SUBSTRING(name, %s)
				AS UNSIGNED
			)
		) AS max_number
		FROM `tabPurchase Invoice`
		WHERE name LIKE %s
		""",
		(len(series) + 1, series + "%"),
		as_dict=True
	)

	max_number = int(result[0].max_number or 0)

	print("Max PI :", max_number)

	if max_number == 0:
		print("ERROR: Tidak ditemukan Purchase Invoice.")
		return

	# Cek tabSeries
	existing = frappe.db.sql(
		"""
		SELECT name, current
		FROM `tabSeries`
		WHERE name = %s
		""",
		(series,),
		as_dict=True
	)

	if existing:

		current = int(existing[0].current or 0)

		print("Current tabSeries :", current)

		if max_number > current:

			print("\nUPDATE")
			print("Before :", current)
			print("After  :", max_number)

			frappe.db.sql(
				"""
				UPDATE `tabSeries`
				SET current = %s
				WHERE name = %s
				""",
				(max_number, series)
			)

		else:

			print("\nTIDAK ADA PERUBAHAN")
			print("tabSeries sudah >= nomor PI terbesar.")

	else:

		print("\nCREATE TAB SERIES")
		print("Name    :", series)
		print("Current :", max_number)

		frappe.db.sql(
			"""
			INSERT INTO `tabSeries`
				(name, current)
			VALUES
				(%s, %s)
			""",
			(series, max_number)
		)

	frappe.db.commit()

	# Verifikasi
	check = frappe.db.sql(
		"""
		SELECT name, current
		FROM `tabSeries`
		WHERE name = %s
		""",
		(series,),
		as_dict=True
	)

	print("\n========================================")
	print("AFTER")
	print("========================================")
	print(check)

	print("\nSUCCESS")


import frappe
import openpyxl


@frappe.whitelist()
def update_item_from_excel():
    file_path = "/home/frappe/frappe-bench/sites/bjm_alan.digitalasiasolusindo.com/public/files/Update Data master Item.xlsx"

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active

    updated = 0
    skipped = 0
    not_found = 0
    errors = []

    headers = [
        str(cell.value).strip() if cell.value is not None else ""
        for cell in ws[1]
    ]

    try:
        item_code_idx = headers.index("Item Code")
        item_name_idx = headers.index("Item Name")
        description_idx = headers.index("Description")
        vendor_idx = headers.index("Vendor")
    except ValueError as e:
        frappe.throw(f"Header Excel tidak ditemukan: {e}")

    for row_no, row in enumerate(
        ws.iter_rows(min_row=2, values_only=True),
        start=2
    ):
        try:
            item_code = row[item_code_idx]
            item_name = row[item_name_idx]
            description = row[description_idx]
            vendor = row[vendor_idx]

            if not item_code:
                skipped += 1
                continue

            item_code = str(item_code).strip()
            item_name = str(item_name).strip() if item_name is not None else ""
            description = str(description).strip() if description is not None else ""
            vendor = str(vendor).strip() if vendor is not None else ""

            if not frappe.db.exists("Item", item_code):
                not_found += 1
                print(f"[NOT FOUND] Row {row_no}: {item_code}")
                continue

            frappe.db.set_value(
                "Item",
                item_code,
                {
                    "item_name": item_name,
                    "description": description,
                    "custom_vendor": vendor,
                },
                update_modified=False,
            )

            updated += 1

            print(
                f"[UPDATED] {item_code} | "
                f"item_name={item_name} | "
                f"description={description} | "
                f"vendor={vendor}"
            )

        except Exception as e:
            errors.append({
                "row": row_no,
                "item_code": row[item_code_idx]
                    if len(row) > item_code_idx else None,
                "error": str(e),
            })

            print(
                f"[ERROR] Row {row_no}: "
                f"{row[item_code_idx] if len(row) > item_code_idx else ''} -> {e}"
            )

    frappe.db.commit()

    print("\n==============================")
    print("UPDATE ITEM SELESAI")
    print(f"Updated   : {updated}")
    print(f"Skipped   : {skipped}")
    print(f"Not Found : {not_found}")
    print(f"Errors    : {len(errors)}")
    print("==============================")

    return {
        "updated": updated,
        "skipped": skipped,
        "not_found": not_found,
        "errors": errors,
    }

def debug_ongoing(date="2026-10-03"):
    pos_invoices = frappe.db.sql("""
        SELECT
            si.name,
            si.owner,
            si.is_return,
            si.posting_date,
            si.pos_profile
        FROM `tabPOS Invoice` si
        WHERE
            si.docstatus = 1
            AND si.is_pos = 1
            AND si.posting_date = %(date)s
            AND si.owner = %(owner)s
    """, {
        "date": date,
        "owner": "kasir2bjm@alan.com"
    }, as_dict=1)

    print("\n" + "=" * 120)
    print("DEBUG ONGOING POS")
    print("Tanggal:", date)
    print("User:", "kasir2bjm@alan.com")
    print("=" * 120)

    for inv in pos_invoices:

        closing = frappe.db.sql("""
            SELECT
                ref.parent AS closing_entry,
                pc.docstatus
            FROM `tabPOS Invoice Reference` ref
            INNER JOIN `tabPOS Closing Entry` pc
                ON pc.name = ref.parent
            WHERE
                ref.pos_invoice = %(invoice)s
        """, {
            "invoice": inv.name
        }, as_dict=1)

        print(
            "Invoice:", inv.name,
            "| Owner:", inv.owner,
            "| Return:", inv.is_return,
            "| Profile:", inv.pos_profile,
            "| Closing:", closing
        )

    print("=" * 120)


def debug_invoice_detail(invoice="RBJM26J030002"):
    print("\n" + "=" * 120)
    print("DEBUG POS INVOICE DETAIL")
    print("Invoice:", invoice)
    print("=" * 120)

    inv = frappe.db.get_value(
        "POS Invoice",
        invoice,
        [
            "name",
            "owner",
            "modified_by",
            "posting_date",
            "posting_time",
            "is_pos",
            "is_return",
            "return_against",
            "pos_profile",
            "docstatus",
            "customer",
            "grand_total"
        ],
        as_dict=True
    )

    if not inv:
        print("Invoice tidak ditemukan.")
        return

    print("\n--- POS INVOICE ---")
    for key, value in inv.items():
        print(f"{key:15}: {value}")

    print("\n--- POS PROFILE ---")

    if inv.pos_profile:
        profile = frappe.db.get_value(
            "POS Profile",
            inv.pos_profile,
            [
                "name",
                "disabled",
                "company",
                "warehouse",
                "currency"
            ],
            as_dict=True
        )

        if profile:
            for key, value in profile.items():
                print(f"{key:15}: {value}")

        # User yang terdaftar pada POS Profile
        profile_users = frappe.db.sql("""
            SELECT
                user,
                default
            FROM `tabPOS Profile User`
            WHERE parent = %(profile)s
        """, {
            "profile": inv.pos_profile
        }, as_dict=True)

        print("\nUser di POS Profile:")
        for u in profile_users:
            print(
                "  User:",
                u.user,
                "| Default:",
                u.default
            )
    else:
        print("POS Profile kosong.")

    print("\n--- POS OPENING ENTRY ---")

    openings = frappe.db.sql("""
        SELECT
            poe.name,
            poe.pos_profile,
            poe.user,
            poe.status,
            poe.posting_date,
            poe.period_start_date,
            poe.period_end_date
        FROM `tabPOS Opening Entry` poe
        WHERE
            poe.posting_date = %(date)s
            AND poe.pos_profile = %(profile)s
        ORDER BY poe.creation DESC
    """, {
        "date": inv.posting_date,
        "profile": inv.pos_profile
    }, as_dict=True)

    if openings:
        for opening in openings:
            print(
                "Opening:", opening.name,
                "| User:", opening.user,
                "| Profile:", opening.pos_profile,
                "| Status:", opening.status,
                "| Posting:", opening.posting_date,
                "| Period:", opening.period_start_date,
                "->",
                opening.period_end_date
            )
    else:
        print("Tidak ada POS Opening Entry.")

    print("\n--- PAYMENT ---")

    payments = frappe.db.sql("""
        SELECT
            parent,
            mode_of_payment,
            amount
        FROM `tabSales Invoice Payment`
        WHERE parent = %(invoice)s
    """, {
        "invoice": invoice
    }, as_dict=True)

    if payments:
        for p in payments:
            mop_type = frappe.db.get_value(
                "Mode of Payment",
                p.mode_of_payment,
                "custom_mop_type"
            )

            print(
                "MOP:", p.mode_of_payment,
                "| Type:", mop_type,
                "| Amount:", p.amount
            )
    else:
        print("Tidak ada payment.")

    print("\n--- POS INVOICE REFERENCE / CLOSING ---")

    closing = frappe.db.sql("""
        SELECT
            ref.name AS reference_name,
            ref.parent AS closing_entry,
            pc.docstatus,
            pc.posting_date
        FROM `tabPOS Invoice Reference` ref
        INNER JOIN `tabPOS Closing Entry` pc
            ON pc.name = ref.parent
        WHERE ref.pos_invoice = %(invoice)s
    """, {
        "invoice": invoice
    }, as_dict=True)

    if closing:
        for c in closing:
            print(
                "Reference:", c.reference_name,
                "| Closing:", c.closing_entry,
                "| Docstatus:", c.docstatus,
                "| Date:", c.posting_date
            )
    else:
        print("TIDAK ADA POS CLOSING REFERENCE.")

    print("\n" + "=" * 120)



def debug_purchase_return():
    pi = "PI-BJB4261000104"

    doc = frappe.get_doc("Purchase Invoice", pi)

    print("=" * 100)
    print("DEBUG PURCHASE RETURN")
    print("PI       :", doc.name)
    print("Date     :", doc.posting_date)
    print("Is Return:", doc.is_return)
    print("Warehouse:", doc.items[0].warehouse if doc.items else "-")
    print("=" * 100)

    for row in doc.items:
        item = row.item_code
        warehouse = row.warehouse

        print("\n" + "-" * 100)
        print("ITEM")
        print("Item Code       :", item)
        print("Item Name       :", row.item_name)
        print("Qty             :", row.qty)
        print("Stock Qty       :", row.stock_qty)
        print("Rate            :", row.rate)
        print("Warehouse       :", warehouse)
        print("Allow Zero Val. :", row.allow_zero_valuation_rate)

        # Item master
        item_data = frappe.db.get_value(
            "Item",
            item,
            ["item_name", "valuation_rate"],
            as_dict=True,
        )

        print("\nITEM MASTER")
        print("Valuation Rate  :", item_data.valuation_rate if item_data else None)

        # Bin
        bin_data = frappe.db.get_value(
            "Bin",
            {
                "item_code": item,
                "warehouse": warehouse,
            },
            ["actual_qty", "valuation_rate", "stock_value"],
            as_dict=True,
        )

        print("\nBIN")
        if bin_data:
            print("Actual Qty      :", bin_data.actual_qty)
            print("Valuation Rate  :", bin_data.valuation_rate)
            print("Stock Value     :", bin_data.stock_value)
        else:
            print("BIN TIDAK ADA")

        # Stock Ledger sebelum tanggal retur
        print("\nSTOCK LEDGER SEBELUM RETURN")

        sle = frappe.db.sql(
            """
            SELECT
                posting_date,
                posting_time,
                voucher_type,
                voucher_no,
                actual_qty,
                qty_after_transaction,
                valuation_rate,
                stock_value,
                stock_value_difference
            FROM `tabStock Ledger Entry`
            WHERE item_code = %s
              AND warehouse = %s
              AND is_cancelled = 0
              AND posting_date <= %s
            ORDER BY posting_date DESC, posting_time DESC, creation DESC
            LIMIT 20
            """,
            (item, warehouse, doc.posting_date),
            as_dict=True,
        )

        if not sle:
            print("TIDAK ADA STOCK LEDGER")
        else:
            for x in sle:
                print(
                    f"{x.posting_date} {x.posting_time} | "
                    f"{x.voucher_type} {x.voucher_no} | "
                    f"qty={x.actual_qty} | "
                    f"qty_after={x.qty_after_transaction} | "
                    f"valuation={x.valuation_rate} | "
                    f"stock_value={x.stock_value} | "
                    f"diff={x.stock_value_difference}"
                )

    print("\n" + "=" * 100)
    print("DEBUG SELESAI")
    print("=" * 100)


def debug_purchase_return_status():
    pi = "PI-BJB4261000104"
    item = "0745760859832"
    warehouse = "Toko - B4"

    doc = frappe.get_doc("Purchase Invoice", pi)

    print("=" * 100)
    print("PI STATUS")
    print("Name       :", doc.name)
    print("Docstatus  :", doc.docstatus)
    print("Is Return  :", doc.is_return)
    print("Posting    :", doc.posting_date)
    print("=" * 100)

    print("\nSLE")
    rows = frappe.db.sql("""
        SELECT
            name,
            posting_date,
            posting_time,
            actual_qty,
            qty_after_transaction,
            valuation_rate,
            stock_value,
            stock_value_difference,
            is_cancelled
        FROM `tabStock Ledger Entry`
        WHERE voucher_type = 'Purchase Invoice'
          AND voucher_no = %s
          AND item_code = %s
          AND warehouse = %s
        ORDER BY creation
    """, (pi, item, warehouse), as_dict=True)

    for row in rows:
        print(row)

    print("\nBIN")
    print(
        frappe.db.get_value(
            "Bin",
            {
                "item_code": item,
                "warehouse": warehouse
            },
            ["actual_qty", "valuation_rate", "stock_value"],
            as_dict=True
        )
    )

    print("=" * 100)


def debug_item_all_stock():
    item = "0745760859832"

    rows = frappe.db.sql("""
        SELECT
            posting_date,
            posting_time,
            warehouse,
            voucher_type,
            voucher_no,
            actual_qty,
            qty_after_transaction,
            valuation_rate,
            stock_value,
            stock_value_difference,
            is_cancelled
        FROM `tabStock Ledger Entry`
        WHERE item_code = %s
          AND is_cancelled = 0
        ORDER BY posting_date, posting_time, creation
    """, (item,), as_dict=True)

    print("=" * 120)
    print("ALL STOCK LEDGER")
    print("ITEM:", item)
    print("=" * 120)

    if not rows:
        print("TIDAK ADA STOCK LEDGER")
        return

    for row in rows:
        print(
            f"{row.posting_date} {row.posting_time} | "
            f"{row.warehouse} | "
            f"{row.voucher_type} {row.voucher_no} | "
            f"qty={row.actual_qty} | "
            f"after={row.qty_after_transaction} | "
            f"valuation={row.valuation_rate} | "
            f"value={row.stock_value} | "
            f"diff={row.stock_value_difference} | "
            f"cancelled={row.is_cancelled}"
        )

def debug_item_documents():
    item = "0745760859832"

    print("=" * 100)
    print("DEBUG ITEM DOCUMENTS")
    print("ITEM:", item)
    print("=" * 100)

    # Purchase Invoice
    print("\nPURCHASE INVOICE")
    rows = frappe.db.sql("""
        SELECT
            pi.name,
            pi.posting_date,
            pi.docstatus,
            pi.is_return,
            pii.qty,
            pii.stock_qty,
            pii.rate,
            pii.warehouse
        FROM `tabPurchase Invoice` pi
        INNER JOIN `tabPurchase Invoice Item` pii
            ON pii.parent = pi.name
        WHERE pii.item_code = %s
        ORDER BY pi.posting_date, pi.creation
    """, (item,), as_dict=True)

    for r in rows:
        print(r)

    # Stock Entry
    print("\nSTOCK ENTRY")
    rows = frappe.db.sql("""
        SELECT
            se.name,
            se.posting_date,
            se.docstatus,
            sei.s_warehouse,
            sei.t_warehouse,
            sei.qty,
            sei.basic_rate
        FROM `tabStock Entry` se
        INNER JOIN `tabStock Entry Detail` sei
            ON sei.parent = se.name
        WHERE sei.item_code = %s
        ORDER BY se.posting_date, se.creation
    """, (item,), as_dict=True)

    for r in rows:
        print(r)

    # Stock Reconciliation
    print("\nSTOCK RECONCILIATION")
    rows = frappe.db.sql("""
        SELECT
            sr.name,
            sr.posting_date,
            sr.docstatus,
            sri.warehouse,
            sri.qty,
            sri.valuation_rate
        FROM `tabStock Reconciliation` sr
        INNER JOIN `tabStock Reconciliation Item` sri
            ON sri.parent = sr.name
        WHERE sri.item_code = %s
        ORDER BY sr.posting_date, sr.creation
    """, (item,), as_dict=True)

    for r in rows:
        print(r)

    print("\n" + "=" * 100)


@frappe.whitelist()
def debug_whitelist():
    import inspect
    import frappe

    method = frappe.get_attr(
        "customer_display.customer_display.doctype.stock_movement_inter"
        ".stock_movement_inter.daftar_gudang"
    )

    return {
        "method": str(method),
        "module": method.__module__,
        "name": method.__name__,
        "source": inspect.getsourcefile(method),
        "in_whitelisted": method in frappe.whitelisted,
        "in_guest_methods": method in frappe.guest_methods,
        "whitelisted_count": len(frappe.whitelisted),
    }


def debug_encryption_key():
    import frappe

    key = frappe.conf.get("encryption_key")

    return {
        "has_key": bool(key),
        "length": len(key or ""),
        "type": type(key).__name__,
        "first_4": (key or "")[:4],
        "last_4": (key or "")[-4:],
    }